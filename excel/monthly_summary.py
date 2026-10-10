"""Build «خلاصه مصرفی ماهیانه» Excel (+ PDF inputs) from plant monthly data."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from config import REPORT_DIR, ensure_dirs
from excel.processor import (
    _normalize_category_code,
    _normalize_fa_header,
    _read_raw_excel,
    _normalize_columns,
)
from excel.work_order import (
    CONSUMPTION_LABELS_FA,
    UNKNOWN_GROUP,
    UNKNOWN_LABEL_FA,
    dominant_work_order,
    group_for_work_order,
    tundish_kg_totals_from_items,
)

SUMMARY_SHEET_NAME = "خلاصه مصرفی ماهیانه"
SUMMARY_FILE_NAME = "خلاصه مصرفی ماهیانه.xlsx"

PERSIAN_MONTHS = {
    1: "فروردین",
    2: "اردیبهشت",
    3: "خرداد",
    4: "تیر",
    5: "مرداد",
    6: "شهریور",
    7: "مهر",
    8: "آبان",
    9: "آذر",
    10: "دی",
    11: "بهمن",
    12: "اسفند",
}

LATIN_TO_PERSIAN_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")

DETAIL_COLUMNS_FA = [
    "کد دسته بندی",
    "کد کالا",
    "مقدار",
    "ضریب",
    "سفارش کار",
    "تاریخ",
    "ماه",
    "واحد",
    "شرح",
]

MONTH_SECTION_COLUMNS_FA = [
    "کد دسته بندی",
    "کد کالا",
    "مقدار",
    "ضریب",
    "مصرف کیلوگرم (مقدار×ضریب)",
    "تاریخ",
    "ماه",
    "واحد",
    "شرح",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
SECTION_FILL = PatternFill("solid", fgColor="7030A0")
SECTION_FONT = Font(bold=True, color="FFFFFF")
MONTH_TITLE_FILL = PatternFill("solid", fgColor="2E75B6")
MONTH_TITLE_FONT = Font(bold=True, color="FFFFFF")

# Row colors by سفارش کار — distinguishable in grayscale / B&W print
# اسلب: خاکستری روشن | بلوم: هاشور نقطه‌ای روشن | بیلت: خاکستری تیره | سایر نواحی: هاشور متراکم
SLAB_FILL = PatternFill(fill_type="solid", fgColor="EDEDED")
BLOOM_FILL = PatternFill(fill_type="lightGray", fgColor="000000", bgColor="FFFFFF")
BILLET_FILL = PatternFill(fill_type="solid", fgColor="9A9A9A")
OTHER_AREA_FILL = PatternFill(fill_type="darkGray", fgColor="000000", bgColor="FFFFFF")
GROUP_FILLS = {
    "slab": SLAB_FILL,
    "bloom": BLOOM_FILL,
    "billet": BILLET_FILL,
    UNKNOWN_GROUP: OTHER_AREA_FILL,
}


def fill_for_work_order(work_order: object) -> PatternFill:
    group = group_for_work_order(work_order) or UNKNOWN_GROUP
    return GROUP_FILLS[group]


def fill_for_group_key(key: str) -> PatternFill:
    return GROUP_FILLS.get(key, OTHER_AREA_FILL)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

COL_WIDTHS = {
    "A": 16.0,
    "B": 15.0,
    "C": 12.0,
    "D": 8.0,
    "E": 14.5,
    "F": 8.0,
    "G": 6.0,
    "H": 10.0,
    "I": 42.0,
}


def to_persian_digits(value: object) -> str:
    return str(value).translate(LATIN_TO_PERSIAN_DIGITS)


def _blank(value: object) -> bool:
    if value is None:
        return True
    try:
        if bool(pd.isna(value)):
            return True
    except (TypeError, ValueError):
        pass
    text = str(value).strip()
    return (not text) or text.casefold() in {"nan", "none", "nat"}


def normalize_fa_text(value: object) -> str:
    if _blank(value):
        return ""
    text = str(value).strip()
    text = text.replace("ي", "ی").replace("ى", "ی").replace("ك", "ک")
    return " ".join(text.split())


def _year_only_latin(value: object) -> int | str | None:
    if _blank(value):
        return None
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    digits = "".join(ch for ch in text if ch.isdigit())
    if len(digits) >= 4:
        return int(digits[:4])
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return text


def _month_int(value: object) -> int | None:
    if _blank(value):
        return None
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    # map Persian month names
    for num, name in PERSIAN_MONTHS.items():
        if name in text:
            return num
    try:
        num = int(float(text))
    except (TypeError, ValueError):
        return None
    if 1 <= num <= 12:
        return num
    return None


def _signed_quantity(qty: float, request_return: object) -> float:
    label = normalize_fa_text(request_return)
    if "برگشت" in label:
        return -abs(float(qty))
    # درخواستی or blank → add
    return float(qty)


def _looks_like_plant_detail(columns: Iterable[str]) -> bool:
    norms = {_normalize_fa_header(c) for c in columns}
    needed = {
        _normalize_fa_header("کد کالا"),
        _normalize_fa_header("مقدار"),
        _normalize_fa_header("ماه"),
        _normalize_fa_header("شرح"),
    }
    return needed.issubset(norms)


@dataclass
class MonthlySummaryData:
    """Aggregated item rows + computed totals for Excel/PDF."""

    items: pd.DataFrame
    grand_kg: float
    group_order: list[str] = field(default_factory=list)
    month_sections: list[dict[str, Any]] = field(default_factory=list)
    # kg/count per slab|bloom|billet|unknown from aggregated item rows
    tundish_totals: dict[str, dict[str, float | int]] = field(default_factory=dict)


def _map_cleaned_to_plant_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Adapt cleaned monthly schema (English) toward plant detail columns."""
    out = pd.DataFrame()
    out["کد دسته بندی"] = df["category_code"] if "category_code" in df.columns else None
    if "id" in df.columns:
        out["کد کالا"] = df["id"]
    elif "material_name" in df.columns:
        out["کد کالا"] = df["material_name"].map(
            lambda v: str(v).split(" - ", 1)[0].strip() if not _blank(v) else None
        )
    else:
        out["کد کالا"] = None
    out["مقدار"] = df["quantity"] if "quantity" in df.columns else None
    out["ضریب"] = df["coefficient"] if "coefficient" in df.columns else 1
    out["سفارش کار"] = df["work_order"] if "work_order" in df.columns else None
    out["تاریخ"] = df["date"] if "date" in df.columns else None
    out["ماه"] = df["month"] if "month" in df.columns else None
    out["واحد"] = df["unit"] if "unit" in df.columns else None
    if "request_return" in df.columns:
        out["درخواستی-برگشتی"] = df["request_return"]
    elif "status" in df.columns:
        out["درخواستی-برگشتی"] = df["status"]
    else:
        out["درخواستی-برگشتی"] = "درخواستی"
    if "description" in df.columns:
        out["شرح"] = df["description"]
    elif "material_name" in df.columns:
        out["شرح"] = df["material_name"].map(
            lambda v: (
                str(v).split(" - ", 1)[1].strip()
                if (not _blank(v) and " - " in str(v))
                else ("" if _blank(v) else str(v).strip())
            )
        )
    else:
        out["شرح"] = None
    return out


def load_monthly_detail(path: Path | str) -> pd.DataFrame:
    """Load best-available monthly detail rows (prefer plant raw sheet)."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"فایل مصرف ماهیانه یافت نشد: {path}")

    # Prefer explicit plant detail sheet when present
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            names = list(wb.sheetnames)
        finally:
            wb.close()
    except Exception:  # noqa: BLE001
        names = []

    sheet: str | int | None = None
    if "ریز اطلاعات" in names:
        sheet = "ریز اطلاعات"
    else:
        for name in names:
            if "ریز اطلاعات" in str(name):
                sheet = name
                break

    if sheet is not None:
        raw = pd.read_excel(path, sheet_name=sheet, engine="openpyxl")
    else:
        raw = _read_raw_excel(path)

    if raw.empty:
        raise ValueError("فایل مصرف ماهیانه خالی است.")

    # Normalize headers to Persian plant names when plant-like
    if _looks_like_plant_detail(raw.columns):
        rename = {}
        target_map = {
            _normalize_fa_header("کد دسته بندی"): "کد دسته بندی",
            _normalize_fa_header("کد کالا"): "کد کالا",
            _normalize_fa_header("مقدار"): "مقدار",
            _normalize_fa_header("ضریب"): "ضریب",
            _normalize_fa_header("سفارش کار"): "سفارش کار",
            _normalize_fa_header("تاریخ"): "تاریخ",
            _normalize_fa_header("ماه"): "ماه",
            _normalize_fa_header("واحد"): "واحد",
            _normalize_fa_header("درخواستی-برگشتی"): "درخواستی-برگشتی",
            _normalize_fa_header("شرح"): "شرح",
        }
        for col in raw.columns:
            key = _normalize_fa_header(col)
            if key in target_map:
                rename[col] = target_map[key]
        detail = raw.rename(columns=rename)
    else:
        # Cleaned / canonical English schema
        normalized = _normalize_columns(raw)
        detail = _map_cleaned_to_plant_columns(normalized)

    required = {"کد کالا", "مقدار", "ماه", "شرح"}
    missing = sorted(required - set(detail.columns))
    if missing:
        raise ValueError(
            "ستون‌های لازم برای خلاصه مصرفی موجود نیست: " + "، ".join(missing)
        )
    if "ضریب" not in detail.columns:
        detail["ضریب"] = 1
    if "درخواستی-برگشتی" not in detail.columns:
        detail["درخواستی-برگشتی"] = "درخواستی"
    if "کد دسته بندی" not in detail.columns:
        detail["کد دسته بندی"] = None
    if "سفارش کار" not in detail.columns:
        detail["سفارش کار"] = None
    if "تاریخ" not in detail.columns:
        detail["تاریخ"] = None
    if "واحد" not in detail.columns:
        detail["واحد"] = None
    return detail


def aggregate_monthly_detail(detail: pd.DataFrame) -> MonthlySummaryData:
    """Aggregate by کد کالا + شرح + ماه; درخواستی add / برگشتی subtract.

    After netting, drop aggregated rows whose مقدار is <= 0 so zero/negative
    net consumption never appears in detail, subtotals, month totals, or WO breakdowns.
    """
    work = detail.copy()
    work["شرح"] = work["شرح"].map(normalize_fa_text)
    def _item_code(v: object) -> str:
        if _blank(v):
            return ""
        text = str(v).strip()
        if text.endswith(".0") and text[:-2].replace("-", "").isalnum():
            text = text[:-2]
        return text

    work["کد کالا"] = work["کد کالا"].map(_item_code)
    work["مقدار"] = pd.to_numeric(work["مقدار"], errors="coerce")
    work = work.loc[work["مقدار"].notna()].copy()
    work["ضریب"] = pd.to_numeric(work["ضریب"], errors="coerce").fillna(1.0)
    work["ماه_num"] = work["ماه"].map(_month_int)
    work = work.loc[work["ماه_num"].notna()].copy()
    work["signed_qty"] = [
        _signed_quantity(q, rr)
        for q, rr in zip(work["مقدار"], work["درخواستی-برگشتی"])
    ]
    work["_order"] = range(len(work))

    # Preserve first-appearance order of شرح for main section
    group_order = list(dict.fromkeys(work["شرح"].tolist()))

    rows: list[dict[str, Any]] = []
    for (item_code, desc, month_num), grp in work.groupby(
        ["کد کالا", "شرح", "ماه_num"], sort=False
    ):
        grp_sorted = grp.sort_values("_order")
        first = grp_sorted.iloc[0]
        qty = float(grp_sorted["signed_qty"].sum())
        coeff = float(first["ضریب"]) if not _blank(first["ضریب"]) else 1.0
        cat = first.get("کد دسته بندی")
        cat_norm = _normalize_category_code(cat)
        # Dominant WO by |signed_qty * coeff|; first appearance as tie-break
        wo_weights = [
            abs(float(sq) * (float(c) if not _blank(c) else 1.0))
            for sq, c in zip(grp_sorted["signed_qty"], grp_sorted["ضریب"])
        ]
        wo_value = dominant_work_order(grp_sorted["سفارش کار"].tolist(), wo_weights)
        rows.append(
            {
                "کد دسته بندی": int(cat_norm) if cat_norm and cat_norm.isdigit() else cat_norm,
                "کد کالا": item_code,
                "مقدار": qty,
                "ضریب": coeff,
                "سفارش کار": wo_value,
                "تاریخ": _year_only_latin(first.get("تاریخ")),
                "ماه": int(month_num),
                "واحد": first.get("واحد"),
                "شرح": desc,
                "مصرف_کیلوگرم": qty * coeff,
                "_first_order": int(grp_sorted["_order"].min()),
            }
        )

    items = pd.DataFrame(rows)
    if items.empty:
        return MonthlySummaryData(
            items=items,
            grand_kg=0.0,
            group_order=[],
            month_sections=[],
            tundish_totals=tundish_kg_totals_from_items(items),
        )

    # Returns already netted into مقدار; drop zero/negative net consumption.
    items = items.loc[items["مقدار"] > 0].copy()
    if items.empty:
        return MonthlySummaryData(
            items=items,
            grand_kg=0.0,
            group_order=[],
            month_sections=[],
            tundish_totals=tundish_kg_totals_from_items(items),
        )
    # Keep first-appearance order, but only for descriptions still present.
    remaining = set(items["شرح"].tolist())
    group_order = [d for d in group_order if d in remaining]

    # Sort within each شرح by month; groups by first appearance
    items["_group_rank"] = items["شرح"].map(
        lambda d: group_order.index(d) if d in group_order else 10**9
    )
    items = items.sort_values(["_group_rank", "ماه", "_first_order"]).reset_index(drop=True)
    grand_kg = float(items["مصرف_کیلوگرم"].sum())

    # Month sections: for each year+month, items sorted by شرح
    month_sections: list[dict[str, Any]] = []
    items["_year"] = items["تاریخ"].map(
        lambda v: int(v) if isinstance(v, int) else (_year_only_latin(v) or 0)
    )
    for (year, month), sec in items.groupby(["_year", "ماه"], sort=True):
        sec_sorted = sec.sort_values(["شرح", "کد کالا"]).reset_index(drop=True)
        month_name = PERSIAN_MONTHS.get(int(month), str(month))
        title = f"مصرف {month_name} ماه {to_persian_digits(year)}"
        total_kg = float(sec_sorted["مصرف_کیلوگرم"].sum())
        month_sections.append(
            {
                "year": int(year) if year else None,
                "month": int(month),
                "month_name": month_name,
                "title": title,
                "total_title": f"جمع کل مصرف {month_name} ماه {to_persian_digits(year)}",
                "total_kg": total_kg,
                "rows": sec_sorted,
            }
        )

    tundish_totals = tundish_kg_totals_from_items(items)
    return MonthlySummaryData(
        items=items,
        grand_kg=grand_kg,
        group_order=group_order,
        month_sections=month_sections,
        tundish_totals=tundish_totals,
    )


def _style_row(ws, row_idx: int, fill: PatternFill | None = None, bold: bool = False) -> None:
    for col in range(1, 10):
        cell = ws.cell(row=row_idx, column=col)
        cell.alignment = CENTER
        cell.border = THIN
        if fill is not None:
            cell.fill = fill
        if bold:
            cell.font = Font(bold=True)


@dataclass
class SummarySheetRow:
    """One logical row of the خلاصه مصرفی ماهیانه sheet (Excel and PDF share this)."""

    kind: str
    # Always length 9 for tabular kinds; ignored for blank/title kinds
    values: list[Any] = field(default_factory=list)
    title: str | None = None
    # Visual cue: header|yellow|green|section|month_title|slab|bloom|billet|unknown|wo
    fill_key: str | None = None
    work_order: Any = None
    # Column schema for this tabular row: "detail" | "month"
    columns_kind: str = "detail"


def _empty9() -> list[Any]:
    return [None] * 9


def _detail_values(item: Any) -> list[Any]:
    return [
        item.get("کد دسته بندی"),
        item.get("کد کالا"),
        item.get("مقدار"),
        item.get("ضریب"),
        item.get("سفارش کار"),
        item.get("تاریخ"),
        item.get("ماه"),
        item.get("واحد"),
        item.get("شرح"),
    ]


def _month_values(item: Any) -> list[Any]:
    return [
        item.get("کد دسته بندی"),
        item.get("کد کالا"),
        item.get("مقدار"),
        item.get("ضریب"),
        item.get("مصرف_کیلوگرم"),
        item.get("تاریخ"),
        item.get("ماه"),
        item.get("واحد"),
        item.get("شرح"),
    ]


def _group_label_with_count(key: str, count: int) -> str:
    if key == UNKNOWN_GROUP:
        base = UNKNOWN_LABEL_FA
    else:
        base = CONSUMPTION_LABELS_FA[key]
    return f"{base} ({to_persian_digits(count)} قلم)"


def build_summary_sheet_rows(data: MonthlySummaryData) -> list[SummarySheetRow]:
    """Build the exact logical rows written to Excel (and mirrored in PDF)."""
    rows: list[SummarySheetRow] = []
    rows.append(
        SummarySheetRow(
            kind="header",
            values=list(DETAIL_COLUMNS_FA),
            fill_key="header",
            columns_kind="detail",
        )
    )

    items = data.items
    if items is not None and not items.empty:
        for desc in data.group_order:
            group = items[items["شرح"] == desc]
            if group.empty:
                continue
            for _, item in group.iterrows():
                wo = item.get("سفارش کار")
                rows.append(
                    SummarySheetRow(
                        kind="data",
                        values=_detail_values(item),
                        fill_key="wo",
                        work_order=wo,
                        columns_kind="detail",
                    )
                )
            subtotal_vals = _empty9()
            subtotal_vals[2] = float(group["مقدار"].sum())
            subtotal_vals[8] = f"جمع {desc}"
            rows.append(
                SummarySheetRow(
                    kind="subtotal",
                    values=subtotal_vals,
                    fill_key="yellow",
                    columns_kind="detail",
                )
            )

    grand_vals = _empty9()
    grand_vals[2] = data.grand_kg
    grand_vals[7] = "کیلوگرم"
    grand_vals[8] = "جمع کل مصرفی"
    rows.append(
        SummarySheetRow(
            kind="grand_total",
            values=grand_vals,
            fill_key="green",
            columns_kind="detail",
        )
    )
    rows.append(SummarySheetRow(kind="blank"))

    rows.append(
        SummarySheetRow(
            kind="section_title",
            title="مصرف به تفکیک ماه و سال",
            fill_key="section",
        )
    )

    for section in data.month_sections:
        rows.append(
            SummarySheetRow(
                kind="month_title",
                title=section["title"],
                fill_key="month_title",
            )
        )
        rows.append(
            SummarySheetRow(
                kind="header",
                values=list(MONTH_SECTION_COLUMNS_FA),
                fill_key="header",
                columns_kind="month",
            )
        )
        for _, item in section["rows"].iterrows():
            wo = item.get("سفارش کار")
            rows.append(
                SummarySheetRow(
                    kind="data",
                    values=_month_values(item),
                    fill_key="wo",
                    work_order=wo,
                    columns_kind="month",
                )
            )
        month_total_vals = _empty9()
        month_total_vals[4] = section["total_kg"]
        month_total_vals[7] = "کیلوگرم"
        month_total_vals[8] = section["total_title"]
        rows.append(
            SummarySheetRow(
                kind="month_total",
                values=month_total_vals,
                fill_key="green",
                columns_kind="month",
            )
        )

        month_totals = tundish_kg_totals_from_items(section["rows"])
        for key in ("slab", "bloom", "billet"):
            info = month_totals.get(key) or {"kg": 0.0, "count": 0}
            count = int(info.get("count") or 0)
            kg = float(info.get("kg") or 0)
            if count == 0 and abs(kg) < 1e-9:
                continue
            gvals = _empty9()
            gvals[4] = kg
            gvals[7] = "کیلوگرم"
            gvals[8] = _group_label_with_count(key, count)
            rows.append(
                SummarySheetRow(
                    kind="group_total",
                    values=gvals,
                    fill_key=key,
                    columns_kind="month",
                )
            )
        unk = month_totals.get(UNKNOWN_GROUP) or {"kg": 0.0, "count": 0}
        unk_kg = float(unk.get("kg") or 0)
        unk_count = int(unk.get("count") or 0)
        if unk_count or abs(unk_kg) > 1e-9:
            gvals = _empty9()
            gvals[4] = unk_kg
            gvals[7] = "کیلوگرم"
            gvals[8] = _group_label_with_count(UNKNOWN_GROUP, unk_count)
            rows.append(
                SummarySheetRow(
                    kind="group_total",
                    values=gvals,
                    fill_key=UNKNOWN_GROUP,
                    columns_kind="month",
                )
            )
        rows.append(SummarySheetRow(kind="blank"))

    # Extra blank ≈ two-row gap before tundish section (matches prior workbook)
    rows.append(SummarySheetRow(kind="blank"))
    rows.append(
        SummarySheetRow(
            kind="section_title",
            title="مصرف مواد بر حسب اسلب، بلوم و بیلت",
            fill_key="section",
        )
    )

    totals = data.tundish_totals or tundish_kg_totals_from_items(data.items)
    for key in ("slab", "bloom", "billet"):
        info = totals.get(key) or {"kg": 0.0, "count": 0}
        count = int(info.get("count") or 0)
        kg = float(info.get("kg") or 0)
        gvals = _empty9()
        gvals[2] = kg
        gvals[7] = "کیلوگرم"
        gvals[8] = _group_label_with_count(key, count)
        rows.append(
            SummarySheetRow(
                kind="group_total",
                values=gvals,
                fill_key=key,
                columns_kind="detail",
            )
        )
    unk = totals.get(UNKNOWN_GROUP) or {"kg": 0.0, "count": 0}
    unk_kg = float(unk.get("kg") or 0)
    unk_count = int(unk.get("count") or 0)
    if unk_count or abs(unk_kg) > 1e-9:
        gvals = _empty9()
        gvals[2] = unk_kg
        gvals[7] = "کیلوگرم"
        gvals[8] = _group_label_with_count(UNKNOWN_GROUP, unk_count)
        rows.append(
            SummarySheetRow(
                kind="group_total",
                values=gvals,
                fill_key=UNKNOWN_GROUP,
                columns_kind="detail",
            )
        )
    check_kg = (
        float((totals.get("slab") or {}).get("kg") or 0)
        + float((totals.get("bloom") or {}).get("kg") or 0)
        + float((totals.get("billet") or {}).get("kg") or 0)
        + unk_kg
    )
    check_vals = _empty9()
    check_vals[2] = check_kg
    check_vals[7] = "کیلوگرم"
    check_vals[8] = (
        "جمع کنترل (اسلب+بلوم+بیلت+سایر نواحی) — باید برابر جمع کل مصرفی باشد"
    )
    rows.append(
        SummarySheetRow(
            kind="check",
            values=check_vals,
            fill_key="green",
            columns_kind="detail",
        )
    )
    return rows


def _fill_for_sheet_row(row: SummarySheetRow) -> PatternFill | None:
    key = row.fill_key
    if key == "header":
        return HEADER_FILL
    if key == "yellow":
        return YELLOW_FILL
    if key == "green":
        return GREEN_FILL
    if key == "section":
        return SECTION_FILL
    if key == "month_title":
        return MONTH_TITLE_FILL
    if key in GROUP_FILLS:
        return GROUP_FILLS[key]
    if key == "wo":
        return fill_for_work_order(row.work_order)
    return None


def build_monthly_summary_workbook(data: MonthlySummaryData) -> Workbook:
    """Create RTL formatted workbook matching the approved sample layout."""
    wb = Workbook()
    ws = wb.active
    ws.title = SUMMARY_SHEET_NAME
    ws.sheet_view.rightToLeft = True

    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    row_idx = 1
    for srow in build_summary_sheet_rows(data):
        if srow.kind == "blank":
            row_idx += 1
            continue

        if srow.kind in {"section_title", "month_title"}:
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=9)
            cell = ws.cell(row=row_idx, column=1, value=srow.title)
            fill = _fill_for_sheet_row(srow)
            font = SECTION_FONT if srow.fill_key == "section" else MONTH_TITLE_FONT
            cell.fill = fill or SECTION_FILL
            cell.font = font
            cell.alignment = CENTER
            for col in range(1, 10):
                c = ws.cell(row=row_idx, column=col)
                c.border = THIN
                if fill is not None:
                    c.fill = fill
            ws.row_dimensions[row_idx].height = 32
            row_idx += 1
            continue

        if srow.kind == "header":
            for col_idx, title in enumerate(srow.values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=title)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = CENTER
                cell.border = THIN
            ws.row_dimensions[row_idx].height = 32
            row_idx += 1
            continue

        # data / subtotal / grand_total / month_total / group_total / check
        for col_idx, val in enumerate(srow.values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)
        fill = _fill_for_sheet_row(srow)
        bold = srow.kind in {"subtotal", "grand_total", "month_total", "group_total", "check"}
        _style_row(ws, row_idx, fill=fill, bold=bold)
        height = 36 if srow.kind == "check" else 32
        ws.row_dimensions[row_idx].height = height
        row_idx += 1

    return wb


def _row_year(value: object) -> int:
    """Extract a positive Jalali year from تاریخ, else 0."""
    if isinstance(value, int) and value > 0:
        # Already year-only (e.g. 1405) or YYYYMMDD int — take leading 4 if long
        if value >= 10000:
            return int(str(value)[:4])
        return value
    parsed = _year_only_latin(value)
    if isinstance(parsed, int) and parsed > 0:
        return parsed
    try:
        num = int(parsed)  # type: ignore[arg-type]
        return num if num > 0 else 0
    except (TypeError, ValueError):
        return 0


def _iter_year_months(
    start: tuple[int, int],
    end: tuple[int, int],
) -> list[tuple[int, int]]:
    """Inclusive (year, month) pairs from start through end."""
    y, m = int(start[0]), int(start[1])
    ey, em = int(end[0]), int(end[1])
    out: list[tuple[int, int]] = []
    while y * 12 + m <= ey * 12 + em:
        out.append((y, m))
        m += 1
        if m > 12:
            m = 1
            y += 1
        if len(out) > 240:
            break
    return out


def filter_summary_by_month_range(
    data: MonthlySummaryData,
    start: tuple[int, int] | None = None,
    end: tuple[int, int] | None = None,
) -> MonthlySummaryData:
    """Keep aggregated months in [start_year/start_month .. end] inclusive; recalculate totals.

    Point-in-time empty filter (no start/end) returns ``data`` unchanged.

    When ``تاریخ`` / year is missing or 0 on aggregated rows, do **not** drop the
    whole file: fall back to matching by ``ماه`` against months that overlap the
    requested range. Display/section year then uses the year from the range
    (single-year span, else end year / last matching pair). Real years are
    preserved when present.
    """
    if start is None and end is None:
        return data
    if data.items is None or data.items.empty:
        return MonthlySummaryData(
            items=data.items.copy() if data.items is not None else pd.DataFrame(),
            grand_kg=0.0,
            group_order=[],
            month_sections=[],
            tundish_totals=tundish_kg_totals_from_items(pd.DataFrame()),
        )

    sy, sm = start if start is not None else (0, 1)
    ey, em = end if end is not None else (9999, 12)
    lo, hi = sy * 12 + sm, ey * 12 + em
    range_pairs = _iter_year_months((sy, sm), (ey, em))
    allowed_months = {m for _, m in range_pairs}
    # Later pairs overwrite → prefer end-of-range year when a month repeats
    month_to_year = {m: y for y, m in range_pairs}

    items = data.items.copy()
    years = items["تاریخ"].map(_row_year)
    months = pd.to_numeric(items["ماه"], errors="coerce").fillna(0).astype(int)

    if bool((years > 0).any()):
        keys = years * 12 + months
        mask = (keys >= lo) & (keys <= hi) & (years > 0) & months.between(1, 12)
        # Rows without year still match by month overlap so partial data is kept
        missing = (years <= 0) & months.isin(list(allowed_months)) & months.between(1, 12)
        mask = mask | missing
    else:
        # All years missing/0 — match by ماه against overlapping months only
        mask = months.isin(list(allowed_months)) & months.between(1, 12)

    filtered = items.loc[mask].copy()
    if filtered.empty:
        return MonthlySummaryData(
            items=filtered,
            grand_kg=0.0,
            group_order=[],
            month_sections=[],
            tundish_totals=tundish_kg_totals_from_items(filtered),
        )

    # Patch missing years for display/sections; prefer real years when present
    patched = filtered["تاریخ"].map(_row_year)
    need = patched <= 0
    if bool(need.any()):
        default_year = sy if sy == ey and sy > 0 else (ey if ey < 9999 else sy)
        fallback = [
            month_to_year.get(int(m), default_year if default_year > 0 else 0)
            for m in filtered.loc[need, "ماه"].tolist()
        ]
        patched = patched.copy()
        patched.loc[need] = fallback
    filtered["تاریخ"] = patched

    remaining = set(filtered["شرح"].tolist())
    group_order = [d for d in data.group_order if d in remaining]
    # preserve first-appearance among remaining if group_order emptied unexpectedly
    if not group_order:
        group_order = list(dict.fromkeys(filtered["شرح"].tolist()))

    filtered["_group_rank"] = filtered["شرح"].map(
        lambda d: group_order.index(d) if d in group_order else 10**9
    )
    sort_cols = ["_group_rank", "ماه"]
    if "_first_order" in filtered.columns:
        sort_cols.append("_first_order")
    filtered = filtered.sort_values(sort_cols).reset_index(drop=True)
    grand_kg = float(filtered["مصرف_کیلوگرم"].sum())

    month_sections: list[dict[str, Any]] = []
    filtered["_year"] = filtered["تاریخ"].map(_row_year)
    for (year, month), sec in filtered.groupby(["_year", "ماه"], sort=True):
        sec_sorted = sec.sort_values(["شرح", "کد کالا"]).reset_index(drop=True)
        month_name = PERSIAN_MONTHS.get(int(month), str(month))
        year_disp = int(year) if year else (sy if sy == ey else ey)
        title = f"مصرف {month_name} ماه {to_persian_digits(year_disp)}"
        total_kg = float(sec_sorted["مصرف_کیلوگرم"].sum())
        month_sections.append(
            {
                "year": int(year_disp) if year_disp else None,
                "month": int(month),
                "month_name": month_name,
                "title": title,
                "total_title": f"جمع کل مصرف {month_name} ماه {to_persian_digits(year_disp)}",
                "total_kg": total_kg,
                "rows": sec_sorted,
            }
        )

    return MonthlySummaryData(
        items=filtered,
        grand_kg=grand_kg,
        group_order=group_order,
        month_sections=month_sections,
        tundish_totals=tundish_kg_totals_from_items(filtered),
    )


def build_monthly_summary(
    source_path: Path | str,
    *,
    excel_out: Path | str | None = None,
    start: tuple[int, int] | None = None,
    end: tuple[int, int] | None = None,
) -> tuple[MonthlySummaryData, Path]:
    """Load → aggregate → optional month/year filter → write Excel.

    ``start`` / ``end`` are inclusive Jalali ``(year, month)`` bounds.
    """
    ensure_dirs()
    detail = load_monthly_detail(source_path)
    data = aggregate_monthly_detail(detail)
    if start is not None or end is not None:
        data = filter_summary_by_month_range(data, start=start, end=end)
    if data.items.empty:
        raise ValueError("پس از تجمیع، ردیفی برای خلاصه مصرفی باقی نماند.")

    if excel_out is None:
        excel_out = REPORT_DIR / SUMMARY_FILE_NAME
    excel_out = Path(excel_out)
    excel_out.parent.mkdir(parents=True, exist_ok=True)
    wb = build_monthly_summary_workbook(data)
    from excel.table_style import finalize_workbook

    finalize_workbook(wb)
    wb.save(excel_out)
    return data, excel_out



def summary_sections_for_pdf(data: MonthlySummaryData) -> list[dict[str, Any]]:
    """Flatten summary into PDF sections that mirror the Excel sheet rows."""
    sheet_rows = build_summary_sheet_rows(data)
    sections: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    def _flush() -> None:
        nonlocal current
        if current is not None:
            sections.append(current)
            current = None

    def _start_table(columns_kind: str, title: str | None, kind: str) -> None:
        nonlocal current
        _flush()
        cols = (
            list(DETAIL_COLUMNS_FA)
            if columns_kind == "detail"
            else list(MONTH_SECTION_COLUMNS_FA)
        )
        current = {
            "title": title,
            "kind": kind,
            "columns": cols,
            "columns_kind": columns_kind,
            "rows": [],
            "grand_kg": data.grand_kg,
        }

    pending_title: str | None = None
    pending_kind: str = "block"

    for srow in sheet_rows:
        if srow.kind == "blank":
            _flush()
            pending_title = None
            continue

        if srow.kind == "section_title":
            _flush()
            sections.append(
                {
                    "title": srow.title,
                    "kind": "banner",
                    "columns": [],
                    "rows": [],
                }
            )
            pending_title = None
            continue

        if srow.kind == "month_title":
            _flush()
            pending_title = srow.title
            pending_kind = "month"
            continue

        if srow.kind == "header":
            # Start a new table; use pending month title if any
            title = pending_title
            kind = pending_kind if pending_title else (
                "main" if srow.columns_kind == "detail" else "month"
            )
            if title is None and srow.columns_kind == "detail" and not sections:
                title = SUMMARY_SHEET_NAME
                kind = "main"
            _start_table(srow.columns_kind, title, kind)
            pending_title = None
            pending_kind = "block"
            continue

        # Tabular content rows — ensure a table exists (tundish has no header in Excel)
        if current is None:
            # After tundish banner: open detail-col table, no title repeat, no header row
            kind = "tundish_wo"
            _start_table(srow.columns_kind or "detail", None, kind)
            current["show_header"] = False

        assert current is not None
        row_dict = {
            "_values": list(srow.values),
            "_kind": srow.kind,
            "_fill_key": srow.fill_key,
            "_work_order": srow.work_order,
        }
        # Also expose by Persian header name for any legacy readers
        cols = current["columns"]
        for i, col_name in enumerate(cols):
            row_dict[col_name] = srow.values[i] if i < len(srow.values) else None
        if srow.kind == "subtotal":
            row_dict["_subtotal"] = True
        current["rows"].append(row_dict)

    _flush()
    return sections
