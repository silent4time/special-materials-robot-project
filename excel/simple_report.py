"""Shared simple Excel (.xlsx) report generator.

Mirrors ``pdf.generator.generate_simple_report_pdf`` inputs so bot and web
share one code path for tabular reports (PDF + Excel together).
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import REPORT_DIR, ensure_dirs
from pdf.generator import DISPLAY_COLUMNS, HEADER_FA, SECTION_ORDER, SIMPLE_HEADER_FA


def _safe_sheet_title(name: str | None, fallback: str = "گزارش", used: set[str] | None = None) -> str:
    """Excel sheet title: max 31 chars, no \\ / * ? : [ ]."""
    raw = (name or fallback).strip() or fallback
    cleaned = re.sub(r'[\\/*?:\[\]]', "-", raw)
    cleaned = cleaned[:31] or fallback[:31]
    if used is None:
        return cleaned
    base = cleaned
    n = 2
    while cleaned in used:
        suffix = f"_{n}"
        cleaned = (base[: 31 - len(suffix)] + suffix) if len(base) + len(suffix) > 31 else base + suffix
        n += 1
    used.add(cleaned)
    return cleaned


def _header_label(col: str, header_map: dict[str, str] | None) -> str:
    labels = header_map or SIMPLE_HEADER_FA
    return str(labels.get(col, col))


def _cell_value(val: Any, col: str | None = None) -> Any:
    if val is None:
        return ""
    try:
        import math

        if isinstance(val, float) and (math.isnan(val) or math.isinf(val)):
            if math.isinf(val):
                return "∞"
            return ""
    except Exception:  # noqa: BLE001
        pass
    # pandas NaN
    try:
        import pandas as pd

        if bool(pd.isna(val)):
            return ""
    except Exception:  # noqa: BLE001
        pass
    return val


def _write_table_sheet(
    ws,
    columns: list[str],
    rows: list[dict[str, Any]],
    *,
    header_map: dict[str, str] | None = None,
    title: str | None = None,
    subtitle: str | None = None,
) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row_idx = 1
    if title:
        ws.cell(row=row_idx, column=1, value=title)
        ws.cell(row=row_idx, column=1).font = Font(bold=True, size=14)
        row_idx += 1
    if subtitle:
        ws.cell(row=row_idx, column=1, value=subtitle)
        row_idx += 1
    if title or subtitle:
        row_idx += 1  # blank spacer

    if not columns:
        ws.cell(row=row_idx, column=1, value="هیچ ستونی تعریف نشده.")
        return

    for c_i, col in enumerate(columns, start=1):
        cell = ws.cell(row=row_idx, column=c_i, value=_header_label(col, header_map))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
    row_idx += 1

    header_row_idx = row_idx - 1  # last written header row
    for row in rows:
        for c_i, col in enumerate(columns, start=1):
            cell = ws.cell(
                row=row_idx,
                column=c_i,
                value=_cell_value(row.get(col, ""), col),
            )
            cell.alignment = center
        row_idx += 1

    # Shared auto-size / wrap / usage_location coloring
    from excel.table_style import apply_product_table_style

    apply_product_table_style(ws, header_row=header_row_idx)


def generate_simple_report_xlsx(
    title: str,
    *,
    subtitle: str | None = None,
    sections: list[dict[str, Any]] | None = None,
    columns: list[str] | None = None,
    rows: list[dict[str, Any]] | None = None,
    empty_message: str = "داده‌ای یافت نشد.",
    output_path: Path | str | None = None,
    header_map: dict[str, str] | None = None,
    filename_stem: str = "simple_report",
) -> Path:
    """Build an .xlsx with the same title/columns/rows (or sections) as the PDF helper.

    One sheet per section (titled safely), or a single sheet named from ``title``.
    Empty sections still get a sheet with ``empty_message`` when columns exist but
    no rows — callers that want «no empty Excel» should skip calling this.
    """
    ensure_dirs()
    if output_path is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = REPORT_DIR / f"{filename_stem}_{stamp}.xlsx"
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".xlsx":
        output_path = output_path.with_suffix(".xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    built: list[dict[str, Any]] = []
    if sections:
        built = list(sections)
    elif columns is not None:
        built = [
            {
                "title": None,
                "columns": list(columns),
                "rows": list(rows or []),
                "empty_message": empty_message,
            }
        ]
    else:
        built = [
            {
                "title": None,
                "columns": [],
                "rows": [],
                "empty_message": empty_message,
            }
        ]

    wb = Workbook()
    # Remove default sheet; recreate per section
    default = wb.active
    wb.remove(default)

    used_titles: set[str] = set()
    for i, section in enumerate(built):
        sec_title = section.get("title")
        sheet_label = section.get("sheet_title") or sec_title
        sheet_name = _safe_sheet_title(
            str(sheet_label) if sheet_label else (title if len(built) == 1 else f"بخش_{i + 1}"),
            fallback="گزارش",
            used=used_titles,
        )
        ws = wb.create_sheet(title=sheet_name)
        cols = list(section.get("columns") or [])
        sec_rows = list(section.get("rows") or [])
        sec_empty = section.get("empty_message") or empty_message
        if not cols or not sec_rows:
            ws.cell(row=1, column=1, value=str(sec_empty))
            continue
        _write_table_sheet(
            ws,
            cols,
            sec_rows,
            header_map=header_map,
            title=title if i == 0 else (str(sec_title) if sec_title else None),
            subtitle=subtitle if i == 0 else None,
        )

    if not wb.sheetnames:
        ws = wb.create_sheet(title=_safe_sheet_title(title))
        ws.cell(row=1, column=1, value=empty_message)

    wb.save(output_path)
    return output_path


def _df_to_row_dicts(df, columns: list[str]) -> list[dict[str, Any]]:
    if df is None:
        return []
    try:
        import pandas as pd

        if not isinstance(df, pd.DataFrame) or df.empty:
            return []
        work = df.copy()
        for c in columns:
            if c not in work.columns:
                work[c] = None
        return work[columns].to_dict(orient="records")
    except Exception:  # noqa: BLE001
        return []


# Analytics sheet specs matching pdf.generator._append_analytics
_ANALYTICS_SHEETS: list[tuple[str, str, list[str]]] = [
    ("مصرف روزانه", "daily_rates", [
        "material_name", "tundish_type", "tundish_id", "avg_daily",
        "total_qty", "days_span", "unit", "source",
    ]),
    ("مصرف بازه", "period_consumption", [
        "material_name", "tundish_type", "tundish_id", "quantity", "unit", "start", "end",
    ]),
    ("موجودی باقیمانده", "remaining", [
        "material_name", "remaining_qty", "unit", "location",
    ]),
    ("مواد بحرانی", "critical", [
        "material_name", "remaining_qty", "avg_daily", "days_of_cover", "unit",
    ]),
    ("پیش‌بینی نیاز", "forecast", [
        "material_name", "tundish_type", "tundish_id", "avg_daily", "days", "forecast_need", "unit",
    ]),
    ("پیشنهاد درخواست", "suggest", [
        "material_name", "avg_daily", "days", "forecast_need",
        "remaining_qty", "suggest_qty", "unit",
    ]),
]


def generate_analytics_report_xlsx(
    frames: dict[str, Any],
    metas: dict[str, dict],
    *,
    analytics: dict[str, Any] | None = None,
    output_path: Path | str | None = None,
    filename_stem: str = "report_analytics",
) -> Path:
    """Multi-sheet workbook of the same key tables used in ``generate_report`` PDF."""
    ensure_dirs()
    if output_path is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = REPORT_DIR / f"{filename_stem}_{stamp}.xlsx"
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".xlsx":
        output_path = output_path.with_suffix(".xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    sections: list[dict[str, Any]] = []

    # Summary of files
    summary_rows = []
    for key in SECTION_ORDER:
        meta = (metas or {}).get(key)
        if not meta:
            continue
        summary_rows.append({
            "منبع": meta.get("label", key),
            "کل ردیف‌ها": meta.get("total_rows", 0),
            "پس از فیلتر نقش": meta.get("visible_rows", 0),
        })
    if summary_rows:
        sections.append({
            "title": "خلاصه فایل‌ها",
            "columns": ["منبع", "کل ردیف‌ها", "پس از فیلتر نقش"],
            "rows": summary_rows,
        })

    if analytics:
        for sheet_title, key, cols in _ANALYTICS_SHEETS:
            frame = analytics.get(key)
            row_dicts = _df_to_row_dicts(frame, cols)
            if row_dicts:
                sections.append({
                    "title": sheet_title,
                    "columns": cols,
                    "rows": row_dicts,
                })

    for key in SECTION_ORDER:
        if key not in (frames or {}):
            continue
        from config import FILE_TYPES

        label = FILE_TYPES.get(key, {}).get("label_fa", key)
        cols = list(DISPLAY_COLUMNS.get(key) or [])
        row_dicts = _df_to_row_dicts(frames.get(key), cols)
        if row_dicts and cols:
            sections.append({
                "title": label,
                "columns": cols,
                "rows": row_dicts,
            })

    return generate_simple_report_xlsx(
        "گزارش تاندیش / خلاصه داده‌های آپلود‌شده",
        sections=sections or [{
            "title": "گزارش",
            "columns": ["پیام"],
            "rows": [{"پیام": "داده‌ای برای اکسل یافت نشد."}],
        }],
        output_path=output_path,
        header_map={**HEADER_FA, **SIMPLE_HEADER_FA},
        filename_stem=filename_stem,
    )


def export_dataframe_xlsx(
    df,
    output_path: Path | str,
    *,
    columns: list[str] | None = None,
    header_map: dict[str, str] | None = None,
    sheet_name: str = "ریز اطلاعات",
) -> Path:
    """Write a DataFrame to xlsx with optional Persian header rename."""
    import pandas as pd

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    if df is None or not isinstance(df, pd.DataFrame) or df.empty:
        raise ValueError("داده‌ای برای خروجی اکسل نیست.")
    work = df.copy()
    use_cols = list(columns) if columns else [c for c in work.columns]
    for c in use_cols:
        if c not in work.columns:
            work[c] = None
    work = work[use_cols]
    labels = header_map or {}
    rename = {c: labels.get(c, c) for c in use_cols}
    work = work.rename(columns=rename)
    safe = _safe_sheet_title(sheet_name)
    work.to_excel(out, index=False, engine="openpyxl", sheet_name=safe)
    from excel.table_style import style_workbook_path

    style_workbook_path(out, sheet_name=safe)
    return out
