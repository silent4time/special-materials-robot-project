"""Shared Excel table styling for product / usage_location sheets.

Applied to منبع اصلی downloads and simple tabular reports so bot + web
exports look consistent.

Palette (entire data-row fill by محل استفاده / usage_location).
Phrases match the منبع اصلی file (no ZWNJ in ریخته گری).

  اسلب                         — soft blue     #BDD7EE
  بلوم                         — soft green    #C6EFCE
  بیلت                         — soft peach    #FCE4D6
  بلوم / بیلت                  — sand mix      #E7F2D3  (green + peach)
  بلوم/اسلب                    — mint mix      #B7DCCB  (green + blue)
  سطح ریخته گری اسلب           — soft lilac    #E2D5F1
  سطح ریخته گری بلوم           — soft teal     #C6E8E3
  سطح ریخته گری بیلت           — soft rose     #F8D3E0
  سطح ریخته گری اسلب/بلوم      — lilac/teal    #D5DEEE  (lilac + teal)
  سایر نواحی                   — pale steel    #D9E1F2
  ترکیبی                       — soft yellow   #FFF2CC
      («،»-joined labels that are not one of the phrases above)
  (empty/blank)                — light gray    #F2F2F2

ZWNJ forms (سطح ریخته‌گری …) use the same fill as the file phrase.
The retired lone label «سطح ریخته گری» uses the slab casting-floor fill.
"""
from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# Documented row-fill palette (hex without #)
USAGE_LOCATION_PALETTE: dict[str, str] = {
    "اسلب": "BDD7EE",
    "بلوم": "C6EFCE",
    "بیلت": "FCE4D6",
    "بلوم / بیلت": "E7F2D3",
    "بلوم/اسلب": "B7DCCB",
    "سطح ریخته گری اسلب": "E2D5F1",
    "سطح ریخته گری بلوم": "C6E8E3",
    "سطح ریخته گری بیلت": "F8D3E0",
    "سطح ریخته گری اسلب/بلوم": "D5DEEE",
    "سایر نواحی": "D9E1F2",
    "ترکیبی": "FFF2CC",
    "": "F2F2F2",
}

_PRODUCT_HEADERS = frozenset({
    "product_name",
    "material_name",
    "item_code_desc",
    "نام محصول",
    "شرح کالا",
    "محصول",
    "نام کالا",
    "کالا",
})
_USAGE_HEADERS = frozenset({
    "usage_location",
    "محل استفاده",
    "محل‌استفاده",
})

_KNOWN_LOCATION_LABELS = frozenset(
    k for k in USAGE_LOCATION_PALETTE if k not in {"ترکیبی", ""}
)

_MAX_COL_WIDTH = 48.0
_MIN_COL_WIDTH = 8.0
_PRODUCT_COL_WIDTH = 42.0
_PRODUCT_ROW_HEIGHT = 30.0  # ~2 visual lines with wrap
_HEADER_ROW_HEIGHT = 22.0


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _canon_location_key(value: object) -> str:
    """Match file phrases and older ZWNJ labels to the same palette key."""
    text = "" if value is None else str(value)
    text = text.replace("ي", "ی").replace("ى", "ی").replace("ك", "ک")
    text = text.replace("\u200c", " ")
    return " ".join(text.strip().split())


def fill_for_usage_location(value: object) -> PatternFill:
    """Return the PatternFill for a usage_location cell value."""
    if value is None:
        return _fill(USAGE_LOCATION_PALETTE[""])
    text = _canon_location_key(value)
    if not text or text.lower() in {"nan", "none", "nat"}:
        return _fill(USAGE_LOCATION_PALETTE[""])
    if text in {"سطح ریخته گری", "سطح ریخته‌گری"}:
        # Retired single label — same fill as slab casting floor.
        return _fill(USAGE_LOCATION_PALETTE["سطح ریخته گری اسلب"])
    if text in USAGE_LOCATION_PALETTE:
        return _fill(USAGE_LOCATION_PALETTE[text])
    # Multi-label (e.g. «اسلب، بیلت») that is not its own phrase → ترکیبی.
    # Dedicated combined phrases (بلوم / بیلت, بلوم/اسلب, …) already matched.
    parts = [p.strip() for p in text.replace(",", "،").split("،") if p.strip()]
    known_hits = [p for p in parts if p in _KNOWN_LOCATION_LABELS]
    if len(known_hits) >= 2 or (len(parts) >= 2 and known_hits):
        return _fill(USAGE_LOCATION_PALETTE["ترکیبی"])
    if len(known_hits) == 1:
        return _fill(USAGE_LOCATION_PALETTE[known_hits[0]])
    return _fill(USAGE_LOCATION_PALETTE[""])


def _header_map(ws: Worksheet, header_row: int) -> dict[int, str]:
    out: dict[int, str] = {}
    max_col = ws.max_column or 0
    for col in range(1, max_col + 1):
        raw = ws.cell(header_row, col).value
        if raw is None:
            continue
        text = str(raw).strip()
        if text:
            out[col] = text
    return out


def _is_product_header(name: str) -> bool:
    return name in _PRODUCT_HEADERS or name.casefold() in {
        h.casefold() for h in _PRODUCT_HEADERS if h.isascii()
    }


def _is_usage_header(name: str) -> bool:
    return name in _USAGE_HEADERS or name.casefold() in {
        h.casefold() for h in _USAGE_HEADERS if h.isascii()
    }


def _display_len(value: object) -> int:
    if value is None:
        return 0
    text = str(value)
    # Rough width: CJK/Persian count as ~1.2 of Latin for Excel columns
    wide = sum(1 for ch in text if ord(ch) > 0x0600)
    return len(text) + int(wide * 0.2)


def apply_product_table_style(
    ws: Worksheet,
    *,
    header_row: int = 1,
    max_col_width: float = _MAX_COL_WIDTH,
    min_col_width: float = _MIN_COL_WIDTH,
    product_row_height: float = _PRODUCT_ROW_HEIGHT,
    color_by_usage: bool = True,
) -> None:
    """Auto-size columns, wrap product description (~2 lines), color by usage_location.

    Safe no-op on empty sheets. Detects English or Persian header aliases.
    """
    if ws.max_row is None or ws.max_column is None:
        return
    if ws.max_row < header_row:
        return

    headers = _header_map(ws, header_row)
    if not headers:
        return

    product_cols = {c for c, name in headers.items() if _is_product_header(name)}
    usage_cols = {c for c, name in headers.items() if _is_usage_header(name)}
    usage_col = next(iter(sorted(usage_cols)), None)

    # Header style pass (bold already set by callers is fine; ensure wrap)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, _name in headers.items():
        cell = ws.cell(header_row, col)
        cell.alignment = header_align
    ws.row_dimensions[header_row].height = _HEADER_ROW_HEIGHT

    # Auto-size from header + data
    for col in range(1, (ws.max_column or 0) + 1):
        maxlen = 0
        for row in range(header_row, (ws.max_row or 0) + 1):
            maxlen = max(maxlen, _display_len(ws.cell(row, col).value))
        width = max(min_col_width, min(max_col_width, float(maxlen) + 2.0))
        if col in product_cols:
            width = min(width, _PRODUCT_COL_WIDTH)
        ws.column_dimensions[get_column_letter(col)].width = width

    wrap_align = Alignment(wrap_text=True, vertical="center", horizontal="right")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in range(header_row + 1, (ws.max_row or 0) + 1):
        fill = None
        if color_by_usage and usage_col is not None:
            fill = fill_for_usage_location(ws.cell(row, usage_col).value)
        for col in range(1, (ws.max_column or 0) + 1):
            cell = ws.cell(row, col)
            if fill is not None:
                cell.fill = fill
            if col in product_cols:
                cell.alignment = wrap_align
            else:
                cell.alignment = center_align
        if product_cols:
            ws.row_dimensions[row].height = product_row_height


def style_workbook_path(
    path: Any,
    *,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> None:
    """Load an .xlsx path, style the active (or named) sheet, save in place."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active
    apply_product_table_style(ws, header_row=header_row)
    finalize_workbook(wb)
    wb.save(path)


# ---------------------------------------------------------------------------
# Phase 3 item 20: one finalizer for EVERY Excel output (standing rule)
# ---------------------------------------------------------------------------
EXCEL_FONT = "Vazirmatn"
HEADER_FILL_HEX = "1F4E79"
_THIN = None


def _thin_border():
    global _THIN
    if _THIN is None:
        from openpyxl.styles import Border, Side

        side = Side(style="thin", color="A6A6A6")
        _THIN = Border(left=side, right=side, top=side, bottom=side)
    return _THIN


def _first_table_row(ws: Worksheet) -> int | None:
    """Header row = first row with ≥2 non-empty cells (title rows have one)."""
    for r in range(1, min(ws.max_row, 30) + 1):
        vals = [c.value for c in ws[r] if c.value not in (None, "")]
        if len(vals) >= 2:
            return r
    return None


def finalize_worksheet(ws: Worksheet) -> None:
    """RTL, Vazirmatn everywhere, centered cells, thin borders on the table,
    header fill when the writer set none. Existing fills/bold/colors are kept."""
    from copy import copy

    from openpyxl.styles import Alignment as _Al, Font as _Font

    ws.sheet_view.rightToLeft = True
    header = _first_table_row(ws)
    max_col = ws.max_column
    _wrap_rows = set()
    # print: whole table width on one page (A4; landscape for wide tables)
    try:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        if max_col > 8:
            ws.page_setup.orientation = "landscape"
    except Exception:  # noqa: BLE001
        pass
    if header and max_col > 1:
        # title / subtitle lines above the table: merge across the table width so a
        # centered long title is not clipped inside column A
        merged_starts = {(r.min_row, r.min_col) for r in ws.merged_cells.ranges}
        for r in range(1, header):
            vals = [c for c in ws[r] if c.value not in (None, "")]
            if len(vals) == 1 and vals[0].column == 1:
                if (r, 1) not in merged_starts:
                    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max_col)
                width = sum(
                    (ws.column_dimensions[get_column_letter(c)].width or 10) for c in range(1, max_col + 1)
                )
                lines = max(1, -(-len(str(vals[0].value)) // max(int(width * 1.1), 20)))
                size = (vals[0].font.sz or 11) if vals[0].font is not None else 11
                ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, lines * size * 1.6)
                _wrap_rows.add(r)
    border = _thin_border()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=max_col):
        r = row[0].row
        in_table = header is not None and r >= header and any(c.value not in (None, "") for c in row)
        for cell in row:
            f = cell.font
            if f is None or f.name != EXCEL_FONT:
                nf = copy(f) if f is not None else _Font()
                nf.name = EXCEL_FONT
                cell.font = nf
            al = cell.alignment
            if cell.value not in (None, "") or in_table:
                if al is None or al.horizontal != "center" or al.vertical != "center" or (r in _wrap_rows and not al.wrap_text):
                    cell.alignment = _Al(
                        horizontal="center", vertical="center",
                        wrap_text=(bool(al.wrap_text) if al is not None else False) or r in _wrap_rows,
                        text_rotation=al.text_rotation if al is not None else 0,
                    )
            if in_table:
                cell.border = border
        if header is not None and r == header:
            for cell in row:
                if cell.value in (None, ""):
                    continue
                fill = cell.fill
                if fill is None or fill.fill_type is None:
                    cell.fill = PatternFill("solid", fgColor=HEADER_FILL_HEX)
                    nf = copy(cell.font)
                    nf.bold, nf.color = True, "FFFFFF"
                    cell.font = nf
                elif not cell.font.bold:
                    nf = copy(cell.font)
                    nf.bold = True
                    cell.font = nf


def finalize_workbook(wb) -> None:
    for ws in wb.worksheets:
        finalize_worksheet(ws)


def finalize_path(path: Any) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    finalize_workbook(wb)
    wb.save(path)
