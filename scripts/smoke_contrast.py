"""Smoke: audit 1405/08/02 — dark-cell contrast, exact number text, Persian codes/activity."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from bot.activity import ACTION_PHRASES, display_message, format_action_phrase
from bot.handlers import _fmt_exact
from excel.contrast import fix_workbook_contrast, is_dark, luminance
from services.units import display_code, fmt_num


def main() -> None:
    assert is_dark("1F3864") and not is_dark("FFFFFF") and not is_dark("FFF2CC")
    assert luminance("000000") < luminance("808080") < luminance("FFFFFF")
    wb = Workbook(); ws = wb.active
    ws["A1"] = "x"; ws["A1"].fill = PatternFill("solid", fgColor="FF1F3864"); ws["A1"].font = Font(color="FF000000")
    ws["B1"] = "y"; ws["B1"].fill = PatternFill("solid", fgColor="FFFFF2CC"); ws["B1"].font = Font(color="FF000000")
    fix_workbook_contrast(wb)
    assert luminance(ws["A1"].font.color.rgb[-6:]) > 0.5, ws["A1"].font.color.rgb
    assert ws["B1"].font.color.rgb.endswith("000000")
    assert fmt_num(1234567.0) == "1234567" and fmt_num(216695.5) == "216695.5"
    assert "e+" not in fmt_num(1e12) and fmt_num(0.123456) == "0.1235"
    assert _fmt_exact(1234.5) == "1,234.5" and _fmt_exact(7.0) == "7"
    assert display_code("منبع", "monthly") == "مصرف ماهیانه" and display_code("کد", "monthly") == "monthly"
    assert "اکسل" in format_action_phrase("report_monthly_summary_xlsx")
    assert "{" not in format_action_phrase("edit_main_source_record")
    assert format_action_phrase("unknown_key_zz") != "unknown_key_zz"
    row = {"message_fa": "کاربر a با آیدی 1 report_surplus_xlsx", "action_key": "report_surplus_xlsx",
           "display_name": "a", "bale_user_id": "1"}
    assert display_message(row) == "کاربر a با آیدی 1 " + ACTION_PHRASES["report_surplus_xlsx"]
    print("OK smoke_contrast")


if __name__ == "__main__":
    main()
