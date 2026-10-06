#!/usr/bin/env python3
"""Smoke: اقلام بحرانی split (شرکت / پیمانکار) × mode (با نوسازی / بدون نوسازی),
ledger applied ONCE, and A4-portrait PDFs for every generator (bot + web share
the same code path).

Self-contained: temp DB + synthetic منبع اصلی; never touches data/bot.db.
Page-1 PNGs are written to reports/smoke_portrait/ for visual checks.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import pandas as pd  # noqa: E402

PNG_DIR = ROOT / "reports" / "smoke_portrait"

COMPANY_ID = "378112341302R"
CONTRACTOR_ID = "378700009002G"


def _inventory_frame() -> pd.DataFrame:
    zero = {
        "bloom_renovation": 0,
        "bloom_patching": 0,
        "slab_renovation": 0,
        "slab_patching": 0,
        "casting_floor": 0,
    }
    rows = [
        # same category on both sides — must NOT be merged
        dict(category_code="1203", id=COMPANY_ID, product_name="بتن 85 شرکت",
             keyword="بتن 85bt", quantity=1000, priority=1, unit="Kg",
             contractor_or_company="شرکت", billet_renovation=10, billet_patching=10, **zero),
        dict(category_code="1203", id=CONTRACTOR_ID, product_name="بتن پلی 85",
             keyword="بتن پلی 85", quantity=500, priority=1, unit="Kg",
             contractor_or_company="پیمانکار", billet_renovation=5, billet_patching=5, **zero),
        # blank column + «0000» id → contractor (id fallback)
        dict(category_code="1655", id="378800001111A", product_name="روکش کشور",
             keyword="روکش کشور", quantity=300, priority=1, unit="Kg",
             contractor_or_company="", billet_renovation=3, billet_patching=0, **zero),
        # Arabic kaf/yeh variant + spaces → company
        dict(category_code="1450", id="378155551234B", product_name="بتن ملات",
             keyword="بتن ملات", quantity=50, priority=1, unit="Kg",
             contractor_or_company=" شركت ", billet_renovation=2, billet_patching=0, **zero),
        # casting floor only → listed in BOTH modes (cf × total tundishes)
        dict(category_code="1451", id="378166661234D", product_name="بتن 80 گان",
             keyword="بتن 80 گان", quantity=40, priority=1, unit="Kg",
             contractor_or_company="شرکت", billet_renovation=0, billet_patching=0,
             **{**zero, "casting_floor": 1}),
        # no rates → never listed
        dict(category_code="9999", id="378199991234C", product_name="بدون نرخ",
             keyword="سایر", quantity=7, priority=1, unit="No",
             contractor_or_company="شرکت", billet_renovation=0, billet_patching=0, **zero),
    ]
    df = pd.DataFrame(rows)
    df["critical_point"] = ""
    return df


def _assert_portrait(pdf: Path) -> int:
    from pypdf import PdfReader

    reader = PdfReader(str(pdf))
    assert reader.pages, pdf
    for i, page in enumerate(reader.pages):
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        rot = int(page.get("/Rotate") or 0) % 180
        if rot:
            w, h = h, w
        assert w < h, f"{pdf.name} page {i + 1} is landscape ({w:.0f}x{h:.0f})"
        assert abs(w - 595.3) < 2 and abs(h - 841.9) < 2, f"{pdf.name} not A4: {w}x{h}"
    return len(reader.pages)


def _png(pdf: Path, name: str) -> Path | None:
    if not shutil.which("pdftoppm"):
        return None
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    stem = PNG_DIR / name
    subprocess.run(
        ["pdftoppm", "-png", "-r", "70", "-f", "1", "-l", "1", "-singlefile", str(pdf), str(stem)],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return stem.with_suffix(".png")


def test_classify() -> None:
    from analytics.critical_items import (
        SEGMENT_COMPANY as CO,
        SEGMENT_CONTRACTOR as CT,
        classify_contractor_or_company as cls,
    )

    assert cls("شرکت") == CO
    assert cls("پیمانکار") == CT
    assert cls(" پيمانكار ") == CT  # Arabic yeh/kaf
    assert cls("پیمان\u200cکار") == CT  # ZWNJ
    assert cls("شركت") == CO
    assert cls("Contractor") == CT and cls("company") == CO
    assert cls("", CONTRACTOR_ID) == CT and cls(None, COMPANY_ID) == CO
    assert cls("nan", "12") == CO  # unknown → company
    assert cls("فولاد خوزستان", CONTRACTOR_ID) == CT  # free text → id rule
    print("classify OK")


def _setup_db(tmp: Path):
    from db.models import Database

    db = Database(tmp / "crit.db")
    db.upsert_user("901", role="owner", display_name="مالک تست")
    user = db.get_user("901")
    inv_dir = tmp / "uploads"
    inv_dir.mkdir(parents=True, exist_ok=True)
    clean = inv_dir / "product_inventory.xlsx"
    _inventory_frame().to_excel(clean, index=False)
    db.save_extracted("901", None, "product_inventory", str(clean), str(clean), 5)
    db.upsert_monthly_tundish_counts(
        jalali_year=1405, jalali_month=6, count_billet=10, count_bloom=0,
        count_slab=0, updated_by="901",
    )
    # one ledger deduction on the company row: -100
    with db.connect() as conn:
        conn.execute(
            """INSERT INTO inventory_ledger
               (item_key, item_id, item_name, delta, reason, ref_type, ref_id,
                bale_user_id, actor_display_name, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (COMPANY_ID, COMPANY_ID, "بتن 85 شرکت", -100.0, "smoke", "smoke", 1,
             "901", "smoke", datetime.now(timezone.utc).isoformat()),
        )
    return db, user


def test_service_split_and_ledger(tmp: Path) -> dict:
    from analytics.critical_items import contractor_column_summary
    from analytics.frames import load_extract_frame, load_primary_inventory
    from services.critical_items_report import generate_critical_items_files

    db, user = _setup_db(tmp)
    raw = load_extract_frame(db, "product_inventory", user)
    inv = load_primary_inventory(db, user)
    q_raw = float(raw.loc[raw["id"] == COMPANY_ID, "quantity"].iloc[0])
    q_inv = float(inv.loc[inv["id"] == COMPANY_ID, "quantity"].iloc[0])
    assert q_raw == 1000 and q_inv == 900, (q_raw, q_inv)
    summary = {(r["value"], r["segment"]): r["rows"] for r in contractor_column_summary(raw)}
    assert summary.get(("شرکت", "company")) == 3 and summary.get(("پیمانکار", "contractor")) == 1, summary

    res = generate_critical_items_files(
        db, user, jalali_year=1405, jalali_month=6, output_dir=tmp / "out", file_prefix="crit",
    )
    assert res.error is None, res.error
    assert res.reno_mode == "with" and "با نوسازی" in res.titles["company"]
    co, ct = res.frames["company"], res.frames["contractor"]
    co_codes = list(co["کد چهاررقمی"].astype(str))
    ct_codes = list(ct["کد چهاررقمی"].astype(str))
    assert sorted(co_codes) == ["1203", "1450", "1451"], co_codes
    assert sorted(ct_codes) == ["1203", "1655"], ct_codes
    r_co = co.loc[co["کد چهاررقمی"].astype(str) == "1203"].iloc[0]
    r_ct = ct.loc[ct["کد چهاررقمی"].astype(str) == "1203"].iloc[0]
    # company: ledger applied once (1000-100=900, NOT 800); need = 10×20
    assert int(r_co["موجودی"]) == 900, r_co["موجودی"]
    assert int(r_co["نیاز"]) == 200
    # contractor: own stock/rates only
    assert int(r_ct["موجودی"]) == 500 and int(r_ct["نیاز"]) == 100
    assert "شرکت" in res.titles["company"] and "پیمانکار" in res.titles["contractor"]
    assert res.company_pdf and res.contractor_pdf and res.xlsx
    from openpyxl import load_workbook

    def _sheet_text(ws) -> str:
        return " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)

    wb = load_workbook(res.xlsx)
    assert wb.sheetnames == ["اقلام بحرانی — شرکت", "اقلام بحرانی — پیمانکار", "توضیحات"], wb.sheetnames
    assert "با نوسازی" in _sheet_text(wb["اقلام بحرانی — شرکت"])
    assert "«با نوسازی»" in _sheet_text(wb["توضیحات"])
    for p in (res.company_pdf, res.contractor_pdf):
        _assert_portrait(p)
    assert "با_نوسازی" in res.company_pdf.name
    print("service split + ledger-once (با نوسازی) OK", co_codes, ct_codes)

    # «بدون نوسازی»: patching only (+ casting floor); renovation-only items drop
    wo = generate_critical_items_files(
        db, user, jalali_year=1405, jalali_month=6, output_dir=tmp / "out",
        file_prefix="crit", reno_mode="without",
    )
    assert wo.error is None and wo.reno_mode == "without", wo.error
    wco, wct = wo.frames["company"], wo.frames["contractor"]
    wco_codes = sorted(wco["کد چهاررقمی"].astype(str))
    wct_codes = sorted(wct["کد چهاررقمی"].astype(str))
    assert wco_codes == ["1203", "1451"], wco_codes  # 1450 (reno only) dropped
    assert wct_codes == ["1203"], wct_codes  # 1655 (reno only) dropped
    n = lambda df, code: int(df.loc[df["کد چهاررقمی"].astype(str) == code, "نیاز"].iloc[0])  # noqa: E731
    assert n(wco, "1203") == 100 and n(wct, "1203") == 50  # 10 × patching
    assert n(wco, "1451") == n(co, "1451") == 10  # casting floor same in both modes
    assert int(wco.loc[wco["کد چهاررقمی"].astype(str) == "1203", "موجودی"].iloc[0]) == 900
    assert "بدون نوسازی" in wo.titles["company"] and "بدون نوسازی" in wo.titles["contractor"]
    wb2 = load_workbook(wo.xlsx)
    assert "بدون نوسازی" in _sheet_text(wb2["اقلام بحرانی — پیمانکار"])
    assert "«بدون نوسازی»" in _sheet_text(wb2["توضیحات"])
    for p in (wo.company_pdf, wo.contractor_pdf):
        _assert_portrait(p)
    assert wo.company_pdf != res.company_pdf and wo.xlsx != res.xlsx
    # pure unit: need formula per mode
    from analytics.critical_items import TundishMonthCounts, monthly_need_for_rates

    cnt = TundishMonthCounts(1405, 6, 70, 0, 100)
    kw = dict(billet_renovation=200, billet_patching=50, bloom_renovation=9, bloom_patching=9,
              slab_renovation=170, slab_patching=30, casting_floor=1, counts=cnt)
    assert monthly_need_for_rates(**kw) == 70 * 250 + 100 * 200 + 170
    assert monthly_need_for_rates(**kw, reno_mode="without") == 70 * 50 + 100 * 30 + 170
    print("service بدون نوسازی OK", wco_codes, wct_codes)
    return {
        "critical_company": res.company_pdf,
        "critical_contractor": res.contractor_pdf,
        "critical_company_no_reno": wo.company_pdf,
        "critical_contractor_no_reno": wo.contractor_pdf,
    }


def test_bot_flow(tmp: Path) -> None:
    """Bot: «تولید گزارش اقلام بحرانی» → company PDF, contractor PDF, xlsx."""
    import config
    import bot.handlers as handlers_mod
    from bot import keyboards as kb

    db, user = _setup_db(tmp)
    handlers_mod.REPORT_DIR = tmp / "reports"
    config.REPORT_DIR = tmp / "reports"

    class FakeClient:
        def __init__(self) -> None:
            self.docs: list[tuple[Path, str]] = []
            self.msgs: list[str] = []

        def send_document(self, chat_id, path, caption=None, **_kw):
            self.docs.append((Path(path), caption or ""))
            return {"ok": True}

        def send_message(self, chat_id, text, reply_markup=None, **_kw):
            self.msgs.append(text)
            return {"ok": True}

        def __getattr__(self, name):  # any other API call → no-op
            return lambda *a, **k: {"ok": True}

    client = FakeClient()
    app = handlers_mod.BotApp(client, db)
    markups: list = []
    orig_send = client.send_message

    def send_message(chat_id, text, reply_markup=None, **kw):
        markups.append(reply_markup)
        return orig_send(chat_id, text, reply_markup=reply_markup, **kw)

    client.send_message = send_message
    msg = {"chat": {"id": 901}, "from": {"id": 901, "first_name": "مالک"}, "text": ""}

    def say(text: str) -> None:
        m = dict(msg, text=text)
        assert app.on_critical_flow_text(m, text), text

    # 1) «با نوسازی» via inline callback
    app.on_critical_report_start(dict(msg, text=kb.BTN_CRITICAL_REPORT))
    say("1405")
    say("شهریور")
    assert "با نوسازی" in client.msgs[-1] and "بدون نوسازی" in client.msgs[-1]
    inline = markups[-1]["inline_keyboard"][0]
    datas = [b["callback_data"] for b in inline]
    assert datas == ["ci|reno|with", "ci|reno|without"], datas
    assert not client.docs  # nothing generated before the choice
    app.handle_callback_query({"id": "cq1", "data": datas[0], "from": msg["from"],
                               "message": {"message_id": 55, "chat": msg["chat"]}})
    names = [p.name for p, _ in client.docs]
    assert names == ["لیست_اقلام_بحرانی_با_نوسازی_شرکت.pdf", "لیست_اقلام_بحرانی_با_نوسازی_پیمانکار.pdf",
                     "لیست_اقلام_بحرانی_با_نوسازی.xlsx"], names
    assert "گزارش اصلی" in client.docs[0][1] and "پیمانکار" in client.docs[1][1]
    assert "با نوسازی" in client.docs[0][1] and "با نوسازی" in client.docs[2][1]
    assert "حالت «با نوسازی»" in client.msgs[-1], client.msgs[-1]
    assert "شرکت (گزارش اصلی): 3 قلم" in client.msgs[-1], client.msgs[-1]
    for p, _ in client.docs[:2]:
        _assert_portrait(p)
    # stale callback after the report → friendly alert, no new docs
    n_docs = len(client.docs)
    app.handle_callback_query({"id": "cq2", "data": datas[1], "from": msg["from"],
                               "message": {"message_id": 55, "chat": msg["chat"]}})
    assert len(client.docs) == n_docs

    # 2) «بدون نوسازی» via typed label
    client.docs.clear()
    app.on_critical_report_start(dict(msg, text=kb.BTN_CRITICAL_REPORT))
    say("1405")
    say("شهریور")
    say("چیز دیگر")  # invalid → re-ask
    assert not client.docs and "یکی از دو دکمه" in client.msgs[-1]
    say(kb.BTN_CRITICAL_RENO_WITHOUT)
    names = [p.name for p, _ in client.docs]
    assert names == ["لیست_اقلام_بحرانی_بدون_نوسازی_شرکت.pdf", "لیست_اقلام_بحرانی_بدون_نوسازی_پیمانکار.pdf",
                     "لیست_اقلام_بحرانی_بدون_نوسازی.xlsx"], names
    assert "حالت «بدون نوسازی»" in client.msgs[-1]
    assert "شرکت (گزارش اصلی): 2 قلم" in client.msgs[-1] and "پیمانکار (گزارش جداگانه): 1 قلم" in client.msgs[-1], client.msgs[-1]
    print("bot critical flow (inline با نوسازی + typed بدون نوسازی) OK")


def test_web(tmp: Path) -> dict:
    from fastapi.testclient import TestClient

    import web.deps as deps
    import web.services.reports as web_reports
    from web.app import create_app

    db, _user = _setup_db(tmp)
    web_reports.REPORT_DIR = tmp / "web_reports"
    import services.critical_items_report as cir

    cir.REPORT_DIR = tmp / "web_reports"
    deps._db = db
    app = create_app()
    app.dependency_overrides[deps.current_user_optional] = lambda: db.get_user("901")
    c = TestClient(app)
    out: dict[str, Path] = {}
    pages = ["/home", "/stock", "/materials/request", "/materials/return", "/reports",
             "/tundish-report", "/tundish-report/settings", "/reports/main-goal"]
    for p in pages:
        r = c.get(p)
        assert r.status_code == 200, (p, r.status_code, r.text[:300])
    r = c.get("/reports")
    assert "اقلام شرکت (PDF)" in r.text and "اقلام پیمانکار (PDF)" in r.text, r.text[-2000:]
    assert 'name="renovation" value="with"' in r.text and 'name="renovation" value="without"' in r.text
    r = c.get("/reports?renovation=without")
    assert 'value="without" checked' in r.text
    import io

    from openpyxl import load_workbook

    for reno in ("with", "without"):
        for seg in ("company", "contractor"):
            r = c.get(
                f"/reports/critical-items?jalali_year=1405&jalali_month=6&segment={seg}&renovation={reno}"
            )
            assert r.status_code == 200 and r.content[:4] == b"%PDF", (seg, reno, r.status_code)
            fname = r.headers.get("content-disposition", "")
            assert ("%D8%A8%D8%AF%D9%88%D9%86" in fname) == (reno == "without"), fname  # «بدون»
            pth = tmp / f"web_critical_{seg}_{reno}.pdf"
            pth.write_bytes(r.content)
            _assert_portrait(pth)
            out[f"web_critical_{seg}_{reno}"] = pth
        r = c.get(f"/reports/critical-items.xlsx?jalali_year=1405&jalali_month=6&renovation={reno}")
        assert r.status_code == 200 and r.content[:2] == b"PK"
        wb = load_workbook(io.BytesIO(r.content))
        label = "بدون نوسازی" if reno == "without" else "با نوسازی"
        notes = " ".join(str(x.value) for row in wb["توضیحات"].iter_rows() for x in row if x.value)
        assert f"«{label}»" in notes, notes[:300]
    r = c.get("/reports/critical-items?jalali_year=1405&jalali_month=7", follow_redirects=False)
    assert r.status_code == 303
    # other web PDF endpoints: PDF (portrait) or a redirect with a Persian error
    for ep in ("/reports/remaining-critical", "/reports/surplus", "/reports/user-activity",
               "/reports/monthly-summary"):
        r = c.get(ep, follow_redirects=False)
        assert r.status_code in (200, 303), (ep, r.status_code)
        if r.status_code == 200 and r.content[:4] == b"%PDF":
            pth = tmp / f"web_{ep.rsplit('/', 1)[-1]}.pdf"
            pth.write_bytes(r.content)
            _assert_portrait(pth)
            out[f"web_{ep.rsplit('/', 1)[-1]}"] = pth
    print("web pages + critical segments OK", sorted(out))
    return out


def test_all_generators(tmp: Path) -> dict:
    """Every PDF generator → A4 portrait; wide 20-column table still fits."""
    from analytics.tundish import daily_rates, remaining
    from pdf.generator import (
        DISPLAY_COLUMNS,
        HEADER_FA,
        generate_monthly_summary_pdf,
        generate_report,
        generate_simple_report_pdf,
        rtl_wrap_lines,
        FONT_NAME,
        _register_fonts,
    )

    _register_fonts()
    lines = rtl_wrap_lines("یک دو سه چهار پنج شش هفت هشت نه ده", 40, FONT_NAME, 8)
    assert len(lines) > 1 and lines[0].startswith("یک"), lines  # logical order kept

    out: dict[str, Path] = {}
    inv = _inventory_frame()
    inv = pd.concat([inv] * 12, ignore_index=True)
    inv["product_name"] = inv["product_name"] + " — شرح طولانی برای آزمون شکستن خط در ستون باریک"
    p = generate_simple_report_pdf(
        "منبع اصلی — آزمون ۲۰ ستون",
        subtitle="عرض صفحه A4 عمودی",
        columns=DISPLAY_COLUMNS["product_inventory"],
        rows=inv.to_dict(orient="records"),
        header_map=HEADER_FA,
        output_path=tmp / "wide20.pdf",
    )
    assert _assert_portrait(p) >= 2
    out["simple_wide20"] = p
    tank = pd.DataFrame({
        "domain": ["تاندیش"] * 3, "tundish_type": ["بیلت"] * 3, "tundish_id": ["T1"] * 3,
        "material_name": ["بتن 85"] * 3, "quantity": [10, 20, 30], "unit": ["Kg"] * 3,
        "assignee_id": ["1"] * 3, "assignee_name": ["الف"] * 3,
        "date": pd.to_datetime(["2026-09-01", "2026-09-02", "2026-09-03"]), "notes": [""] * 3,
    })
    rates = daily_rates(tank, None)
    rem = remaining(inv)
    p = generate_report(
        {"tank_consumption": tank, "product_inventory": inv},
        {"tank_consumption": {"label": "مصرف", "total_rows": 3, "visible_rows": 3},
         "product_inventory": {"label": "منبع اصلی", "total_rows": len(inv), "visible_rows": len(inv)}},
        {"bale_user_id": "901", "role": "owner", "display_name": "مالک"},
        tmp / "general.pdf",
        analytics={"daily_rates": rates, "remaining": rem, "critical": rem.head(0),
                   "forecast": rates.head(0), "suggest": rates.head(0), "period_consumption": tank,
                   "days": 3, "start": "2026-09-01", "end": "2026-09-03"},
    )
    _assert_portrait(p)
    out["general_report"] = p
    p = generate_monthly_summary_pdf(
        [{"kind": "banner", "title": "شهریور ۱۴۰۵"},
         {"title": "بیلت", "columns": ["ماده", "مقدار", "واحد", "توضیح"],
          "rows": [{"ماده": "بتن 85", "مقدار": 1200.5, "واحد": "Kg", "توضیح": "متن نسبتا طولانی برای شکستن خط"},
                   {"_kind": "subtotal", "_values": ["جمع", 1200.5, "Kg", ""]}]}],
        grand_kg=1200.5,
        output_path=tmp / "monthly.pdf",
    )
    _assert_portrait(p)
    out["monthly_summary"] = p
    print("all generators portrait OK", sorted(out))
    return out


def main() -> int:
    print("smoke critical split + portrait…")
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        test_classify()
        pdfs: dict[str, Path] = {}
        for sub in ("svc", "bot", "web", "gen"):
            (tmp / sub).mkdir()
        pdfs.update(test_service_split_and_ledger(tmp / "svc"))
        test_bot_flow(tmp / "bot")
        pdfs.update(test_web(tmp / "web"))
        pdfs.update(test_all_generators(tmp / "gen"))
        pngs = [_png(p, k) for k, p in pdfs.items()]
        print("page-1 PNGs:", ", ".join(str(x.name) for x in pngs if x))
    print("SMOKE_CRITICAL_PORTRAIT_OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
