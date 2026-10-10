#!/usr/bin/env python3
"""Offline smoke: گزارش هدف اصلی — period detect, parse, compute, PDF/xlsx, DB."""
from __future__ import annotations

import _smoke_env  # noqa: F401,E402  — temp DB/reports/uploads before config import

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
import _smoke_isolation  # noqa: E402 - real uploads/ stays untouched

_smoke_isolation.isolate_uploads()


def _write_samples(tmp: Path) -> dict[str, Path]:
    from openpyxl import Workbook

    tmp.mkdir(parents=True, exist_ok=True)
    period = "شهریور ۱۴۰۵"

    # 1) production — CCM columns + PRODUCT row
    prod = tmp / f"آمار_تولید_{period}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "تولید"
    ws["A1"] = f"آمار تولید {period}"
    ws.append(["CCM", "CCM1", "CCM2", "CCM3", "CCM4", "CCM5"])
    ws.append(["PRODUCT (TON)", 10000, 8000, 12000, 15000, 9000])
    wb.save(prod)

    def cons_file(name: str, section_fa: str, tundish: int, heats: int, materials: list[tuple[str, float]]) -> Path:
        path = tmp / f"مصرف_تاندیش_{section_fa}_{period}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "مصرف"
        ws["A1"] = f"مصرف تاندیش {section_fa} — {period}"
        ws.append(["NO. TUNDISH", tundish])
        ws.append(["NO. HEAT", heats])
        ws.append(["COATING (Kg)", materials[0][1]])
        ws.append(["CASTABLE (Kg)", materials[1][1]])
        ws.append(["SPECIAL BRICK", materials[2][1]])
        ws.append(["KG /TON", 1.2])  # should be skipped as material
        wb.save(path)
        return path

    return {
        "production": prod,
        "billet_consumption": cons_file("billet", "بیلت", 40, 320, [("c", 8000), ("a", 1200), ("b", 3500)]),
        "bloom_consumption": cons_file("bloom", "بلوم", 25, 200, [("c", 5000), ("a", 900), ("b", 2100)]),
        "slab_consumption": cons_file("slab", "اسلب", 30, 250, [("c", 6000), ("a", 1100), ("b", 2800)]),
    }


def test_period_and_compute(tmp: Path) -> None:
    from services import main_goal_report as mg

    files = _write_samples(tmp)
    # period from filenames
    for kind, path in files.items():
        period, src = mg.detect_period(path, filename=path.name)
        assert period is not None, (kind, src)
        assert period.kind == "month" and period.month == 6 and period.year == 1405, period
        print(f"  period {kind}: {period.label_fa()} via {src}")

    # mismatch error
    bad = tmp / "مصرف_تاندیش_بیلت_مهر_۱۴۰۵.xlsx"
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "مهر ۱۴۰۵"
    ws.append(["NO. TUNDISH", 1])
    ws.append(["COATING (Kg)", 10])
    wb.save(bad)
    mismatched = dict(files)
    mismatched["billet_consumption"] = bad
    result = mg.compute_main_goal(mismatched, filenames={k: Path(v).name for k, v in mismatched.items()})
    assert not result.ok
    assert "بازه زمانی فایل‌ها یکسان نیست" in (result.error_fa or ""), result.error_fa
    print("  mismatch error OK")

    result = mg.compute_main_goal(
        files,
        filenames={k: Path(v).name for k, v in files.items()},
        target_tons=100000,
    )
    assert result.ok, result.error_fa
    assert result.production is not None
    # CCM1+CCM2=18000 slab, CCM3=12000 bloom, CCM4+CCM5=24000 billet
    assert abs(result.production.slab_tons - 18000) < 0.1, result.production
    assert abs(result.production.bloom_tons - 12000) < 0.1, result.production
    assert abs(result.production.billet_tons - 24000) < 0.1, result.production
    assert result.consumptions["billet"].tundish_count == 40
    assert result.consumptions["slab"].melt_count == 250
    assert result.rate_rows, "expected material rates"
    # coating billet 8000 / 40 tundish = 200; / 24000 ton
    coating = [r for r in result.rate_rows if "COATING" in r.material_name.upper() or "coating" in r.material_name.casefold()]
    assert coating, [r.material_name for r in result.rate_rows]
    billet_coat = [r for r in coating if r.section == "billet"][0]
    assert abs((billet_coat.per_tundish or 0) - 200) < 0.01
    assert billet_coat.per_ton is not None and billet_coat.per_ton > 0
    assert result.sections
    print("  compute OK", result.summary_text().split("\n")[1:3])


def test_db_and_pdf(tmp: Path) -> None:
    from db.models import Database
    from excel.simple_report import generate_simple_report_xlsx
    from pdf.generator import generate_simple_report_pdf
    from services import main_goal_report as mg
    from bot.jalali import format_date, tehran_now

    db_path = tmp / "mg.db"
    db = Database(str(db_path))
    files = _write_samples(tmp / "f")
    result = mg.compute_main_goal(files, filenames={k: Path(v).name for k, v in files.items()})
    assert result.ok
    now = tehran_now()
    row = db.insert_main_goal_report(
        period_key=result.period.key(),
        period_label=result.period.label_fa(),
        period_json="{}",
        files_json="{}",
        results_json=mg.persist_payload(result, {}),
        summary_text=result.summary_text(),
        target_tons=None,
        source="test",
        bale_user_id="1",
        actor_display_name="تست",
        created_at=mg.utc_now_iso(),
        created_at_tehran=now.isoformat(),
        jalali_date=format_date(now),
    )
    assert row["id"]
    assert db.list_main_goal_reports(limit=1)[0]["period_label"]

    pdf = tmp / "mg.pdf"
    xlsx = tmp / "mg.xlsx"
    generate_simple_report_pdf(
        mg.TITLE_FA,
        subtitle=result.period.label_fa(),
        sections=result.sections,
        output_path=pdf,
        filename_stem="mg_smoke",
    )
    generate_simple_report_xlsx(
        mg.TITLE_FA,
        subtitle=result.period.label_fa(),
        sections=result.sections,
        output_path=xlsx,
        filename_stem="mg_smoke",
    )
    assert pdf.stat().st_size > 500
    assert xlsx.stat().st_size > 500
    print("  db+pdf/xlsx OK", pdf.stat().st_size, xlsx.stat().st_size)


def test_bot_menu_wiring() -> None:
    from bot import keyboards as kb
    from services import main_goal_report as mg

    assert kb.BTN_MAIN_GOAL in str(kb.main_menu({"active": 1, "role": "owner"}))  # main menu since phase 1
    assert kb.BTN_MAIN_GOAL not in str(kb.main_menu({"active": 1, "role": "technician"}))
    assert mg.can_run({"active": 1, "role": "owner"})
    assert mg.can_run({"active": 1, "role": "manager"})
    assert mg.can_run({"active": 1, "role": "responsible_officer"})
    assert not mg.can_run({"active": 1, "role": "technician"})
    print("  menu/roles OK")


def main() -> None:
    print("smoke main_goal_report…")
    test_bot_menu_wiring()
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        test_period_and_compute(tmp)
        test_db_and_pdf(tmp)
    print("ALL OK")


if __name__ == "__main__":
    main()
