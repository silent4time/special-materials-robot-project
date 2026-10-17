"""Standing rule (1405/07/25): every material report table starts with the main-source base columns.

    کد ۴ رقمی، آیدی کد، شرح، کلیدواژه، موجودی، محل استفاده، تأمین‌کننده

Part A (portable): ``services.material_columns`` on a synthetic main source (by id, by
«آیدی - شرح», per-code joined with segment, shared-group rows, fallback, rename, bold-row remap).
Part B (``--db``, default data/bot.db when present; a *copy* is used, never written): generates
every material report (bot handlers with a fake client, services, web, assistant) into the smoke
temp dir, then asserts that every material table header in every xlsx and PDF starts with the base
columns, wide PDFs are landscape, and base «موجودی» equals the main source (+ ledger).
"""
from __future__ import annotations

import _smoke_env  # noqa: F401,E402

import argparse
import os
import shutil
import sys
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

ap = argparse.ArgumentParser()
ap.add_argument("--db", default=str(ROOT / "data" / "bot.db"), help="copy of this DB is used (never written)")
ARGS, _ = ap.parse_known_args()
if ARGS.db and Path(ARGS.db).is_file():
    dst = _smoke_env.SMOKE_TMP / "copy.db"
    for suf in ("", "-wal", "-shm"):
        if Path(ARGS.db + suf).exists():
            shutil.copy2(ARGS.db + suf, str(dst) + suf)
    os.environ["DATABASE_PATH"] = str(dst)
    HAVE_DB = True
else:
    HAVE_DB = False

FAILS: list[str] = []


def check(ok: bool, msg: str) -> None:
    print(("  ✓ " if ok else "  ✗ ") + msg)
    if not ok:
        FAILS.append(msg)


# --- Part A --------------------------------------------------------------------
def part_a() -> None:
    import pandas as pd

    from services import material_columns as mc

    print("== part A: helper ==")
    inv = pd.DataFrame([
        {"category_code": 1672, "id": "378121641252L", "product_name": "استوپر A", "keyword": "استوپر",
         "quantity": 10, "usage_location": "اسلب", "contractor_or_company": "شرکت"},
        {"category_code": 1672, "id": "378100001302R", "product_name": "استوپر B", "keyword": "استوپر",
         "quantity": 5, "usage_location": "بلوم", "contractor_or_company": "پیمانکار"},
        {"category_code": 1674, "id": "378121649999X", "product_name": "نازل C", "keyword": "نازل بیرونی",
         "quantity": 7, "usage_location": "اسلب", "contractor_or_company": "شرکت"},
    ])
    idx = mc.index_from(inv)
    check(len(idx) == 3 and idx.by_identifier("378100001302R")[mc.B_SUP] == "پیمانکار",
          "id rule → پیمانکار (digits 5–8 = 0000)")
    cols, rows, _ = mc.apply_rows(["شناسه", "مقدار"], [{"شناسه": "378121641252L", "مقدار": 3}],
                                  {"id": "شناسه", "mode": "per_id"}, idx)
    check(cols == mc.BASE_COLUMNS + ["مقدار"] and rows[0][mc.B_QTY] == 10 and rows[0][mc.B_DESC] == "استوپر A",
          f"by id {cols}")
    cols, rows, _ = mc.apply_rows(["material_name", "remaining_qty"],
                                  [{"material_name": "378121649999X - نازل C", "remaining_qty": 2}],
                                  mc.SPEC_SITE_REMAINING, idx)
    check(cols[-1] == "موجودی سایت" and rows[0][mc.B_CODE] == "1674" and rows[0]["موجودی سایت"] == 2,
          "by «آیدی - شرح» + rename")
    sec = mc.apply_section({"columns": ["کد چهاررقمی", "ردیف", "کد و شرح کالا", "موجودی", "نیاز"],
                            "rows": [{"کد چهاررقمی": "1672", "ردیف": 1, "موجودی": 10, "نیاز": 4},
                                     {"کد چهاررقمی": "1672 / 1674", "ردیف": "", "موجودی": 17, "نیاز": 9}],
                            "bold_rows": [1],
                            "material": {"code": "کد چهاررقمی", "segment": "company", "mode": "joined"}}, idx)
    r0, r1 = sec["rows"]
    check(sec["columns"] == mc.BASE_COLUMNS + ["ردیف", "نیاز"], f"per-code columns {sec['columns']}")
    check(r0[mc.B_ID] == "378121641252L" and r0[mc.B_QTY] == 10, "per-code: only the segment's IDs")
    check(r1[mc.B_CODE] == "1672 / 1674" and r1[mc.B_QTY] == 17 and not r1[mc.B_ID], "shared-group row")
    check(sec["bold_rows"] == [1], "bold rows kept")
    sec = mc.apply_section({"columns": ["کد", "نام"], "rows": [{"کد": "1672", "نام": "x"}],
                            "material": {"code": "کد", "mode": "joined"}}, idx)
    check(sec["rows"][0][mc.B_ID] == "378121641252L، 378100001302R" and sec["rows"][0][mc.B_QTY] == 15,
          "joined IDs + summed stock")
    sec = mc.apply_section({"columns": ["شناسه", "شرح"], "rows": [{"شناسه": "NOPE1", "شرح": "ناشناخته"}],
                            "material": {"id": "شناسه", "name": "شرح"}}, idx)
    check(sec["rows"][0][mc.B_ID] == "NOPE1" and sec["rows"][0][mc.B_DESC] == "ناشناخته", "fallback")
    check(mc.needs_landscape([{"columns": list(range(9)), "rows": [1]}]), "landscape threshold")


# --- Part B --------------------------------------------------------------------
MARKERS = {"کد ۴ رقمی", "آیدی کد", "شرح", "شرح کالا", "کد و شرح کالا", "شناسه", "شناسه مواد", "شناسه کالا",
           "کد دسته", "کد چهاررقمی", "ماده", "کلید واژه", "کلیدواژه"}
# not main-source items: tundish-table operations («تعویض شرود»…) of the period report
EXEMPT_HEADERS = [["ردیف", "بخش", "ماده", "واحد", "مصرف", "ماه‌ها"]]


def _is_material(header: list[str]) -> bool:
    return bool(set(header) & MARKERS) and set(header) not in [set(h) for h in EXEMPT_HEADERS]


def xlsx_headers(path: Path) -> list[tuple[str, list[str]]]:
    from openpyxl import load_workbook

    out = []
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if len(cells) >= 2 and _is_material(cells) and (set(cells) & MARKERS) and all(isinstance(c, str) for c in row if c is not None):
                out.append((ws.title, cells))
                break
    return out


class PdfSpy:
    def __init__(self) -> None:
        import pdf.generator as g

        self.g = g
        self.headers: list[list[str]] = []
        real = g.build_rtl_table

        def spy(header_texts, rows_texts, **k):
            if header_texts:
                self.headers.append([str(h) for h in header_texts])
            return real(header_texts, rows_texts, **k)

        g.build_rtl_table = spy

    def take(self) -> list[list[str]]:
        h, self.headers = self.headers, []
        return h


def part_b() -> None:
    from pypdf import PdfReader

    from services import material_columns as mc

    print("== part B: every material report (DB copy) ==")
    from analytics.frames import load_primary_inventory
    from db.models import Database

    db = Database()
    owner = next((u for u in db.list_users() if u.get("role") == "owner" and u.get("active")), None)
    if owner is None:
        print("  (no owner — skipped)")
        return
    inv = load_primary_inventory(db, owner)
    idx = mc.index_from(inv)
    spy = PdfSpy()
    out = _smoke_env.SMOKE_TMP / "material_reports"
    out.mkdir()
    reports: dict[str, list[Path]] = {}
    pdf_heads: dict[str, list[list[str]]] = {}

    def collect(name: str, files, *, expect_material: bool = True) -> None:
        files = [Path(f) for f in files if f]
        reports[name] = files
        pdf_heads[name] = spy.take()
        if not files:
            check(False, f"{name}: no files")
            return
        mats_x, mats_p = [], [h for h in pdf_heads[name] if _is_material(h)]
        for f in files:
            if f.suffix == ".xlsx":
                mats_x += xlsx_headers(f)
        bad = [h for _s, h in mats_x if not mc.header_starts_with_base(h)]
        bad_p = [h for h in mats_p if not mc.header_starts_with_base(h)]
        if expect_material:
            check(bool(mats_x) or not any(f.suffix == ".xlsx" for f in files), f"{name}: xlsx has material table")
            check(bool(mats_p) or not any(f.suffix == ".pdf" for f in files), f"{name}: pdf has material table")
        check(not bad, f"{name}: xlsx headers start with base {bad[:1]}")
        check(not bad_p, f"{name}: pdf headers start with base {bad_p[:1]}")
        for f in files:
            if f.suffix == ".pdf":
                wide = any(len(h) >= 9 for h in mats_p)
                p = PdfReader(str(f)).pages[0].mediabox
                if wide:
                    check(float(p.width) > float(p.height), f"{name}: wide table → landscape")

    # services ------------------------------------------------------------------
    from services import critical_items_report as cr
    from services import n_tundish_report as nt
    from services import period_consumption as pc
    from services import comprehensive_report as comp
    from services import inbound_report as ib

    res = cr.generate_critical_items_files(db, owner, output_dir=out / "crit", inventory_df=inv)
    if res.error:
        print("  (critical:", res.error[:80], ")")
    else:
        # TEMPLATE-LOCKED exception (1405/07/25 user decision): اقلام بحرانی follows the
        # user's sample «140506» — NO base material columns, not the shared styler.
        from openpyxl import load_workbook as _lw

        from services.critical_items_report import TEMPLATE_COLUMNS as _RC
        from excel.table_style import template_locked_name

        spy.take()
        _wb = _lw(res.xlsx)
        check(template_locked_name(_wb) == "critical_items", "critical: xlsx template-locked")
        for _ws in _wb.worksheets:
            _hdr = [c.value for c in _ws[2]]
            check(_hdr == list(_RC), f"critical {_ws.title}: sample columns, no base columns {_hdr[:3]}")
            check(mc.B_ID not in _hdr, f"critical {_ws.title}: no «{mc.B_ID}»")
    r = nt.generate_files(db, owner, "slab", 4, output_dir=out / "nt", inventory_df=inv)
    if getattr(r, "pdf", None):
        collect("نیاز N تاندیش", [r.pdf, r.xlsx])
        for row in r.rows:
            seg = mc._seg_key(row.get(nt.COL_SUPPLIER))
            base = sum(x[mc.B_QTY] or 0 for x in idx.for_code(row.get(nt.COL_CODE), seg, planning=True))
            if abs(base - float(row.get(nt.COL_STOCK) or 0)) > 1e-6 and "/" not in str(row.get(nt.COL_CODE)):
                check(False, f"N tundish {row.get(nt.COL_CODE)}: base {base} ≠ {row.get(nt.COL_STOCK)}")
                break
        check(True, "N tundish: base stock = report stock")
    end = date.today()
    start = end - timedelta(days=89)
    p = pc.generate_files(db, start, end, range_label="۹۰ روز", output_dir=out / "pc")
    if getattr(p, "pdf", None):
        collect("مصرف بازه", [p.pdf, p.xlsx], expect_material=bool(p.site_rows))
    rep = ib.latest_report(db)
    if rep:
        a, b = ib.build_report_files(db, rep, out_dir=out / "inb")
        collect("ورودی انبار", [a, b], expect_material=bool(rep.get("lines")))
        pv = ib.preview_table(rep)
        check(pv is None or mc.header_starts_with_base(pv["header"]), "web inbound preview header")
    c = comp.generate_files(db, owner, start, end, range_label="۹۰ روز")
    if c.pdf:
        collect("گزارش جامع", [c.pdf, c.xlsx])
        check(float(PdfReader(str(c.pdf)).pages[0].mediabox.width) > 800, "comprehensive PDF landscape")

    # web -----------------------------------------------------------------------
    from web.services import reports as wr

    for name, fn in (("پوشش کوتاه‌مدت (وب)", wr.generate_remaining_critical_files),
                     ("مازاد (وب)", wr.generate_surplus_files), ("مصرف روزانه (وب)", wr.generate_daily_files)):
        a, b, err = fn(db, owner, days=90)
        if err:
            print(f"  ({name}: {err[:60]})")
            spy.take()
            continue
        collect(name, [a, b])

    # bot handlers (fake client) --------------------------------------------------
    import smoke_phase1 as p1
    from bot.handlers import BotApp

    client = p1.FakeClient()
    app = BotApp(client, db)  # type: ignore[arg-type]
    msg = {"from": {"id": int(owner["bale_user_id"]), "first_name": "x"}, "chat": {"id": int(owner["bale_user_id"])},
           "text": "x"}

    def bot_run(name: str, call, expect_material: bool = True) -> None:
        n0 = len(client.docs)
        try:
            call()
        except Exception as exc:  # noqa: BLE001
            check(False, f"{name}: {exc!r}")
            return
        docs = [Path(d) for _c, d in client.docs[n0:]]
        if not docs:
            print(f"  ({name}: no document — no data)")
            spy.take()
            return
        collect(name, docs, expect_material=expect_material)

    bot_run("مازاد (ربات)", lambda: app._run_surplus_report(msg, owner, start, end, "۹۰ روز"))
    bot_run("پوشش کوتاه‌مدت (ربات)", lambda: app._run_remaining_critical(msg, owner, start, end, "۹۰ روز"))
    bot_run("مصرف روزانه (ربات)", lambda: app._run_daily_report(msg, owner, start, end, "۹۰ روز"))
    import bot.keyboards as kb

    btn = getattr(kb, "BTN_LIST_CATEGORIES", None) or getattr(kb, "BTN_CATEGORY_LIST", None)
    if hasattr(app, "on_list_categories"):
        bot_run("لیست کد دسته‌بندی و موجودی", lambda: app.on_list_categories(msg))
    elif btn:
        bot_run("لیست کد دسته‌بندی و موجودی",
                lambda: app.handle_update({"update_id": 1, "message": dict(msg, text=btn)}))
    try:
        lines, err = app._build_mr_lines(owner, 30)
    except Exception as exc:  # noqa: BLE001
        lines, err = [], repr(exc)
    if lines:
        bot_run("پیش‌نویس درخواست مواد", lambda: app._send_simple_pdf_report(
            msg, title="پیش‌نویس درخواست مواد", columns=["ردیف", "شناسه", "شرح", "گروه", "موجودی", "مصرف روز",
                                                         "پیشنهاد", "واحد"],
            rows=app._mr_lines_to_pdf_rows(lines), material={"id": "شناسه", "name": "شرح", "mode": "per_id"},
            filename_stem="mr_smoke"))
        src = (ROOT / "bot" / "handlers.py").read_text(encoding="utf-8")
        check(src.count('material={"id": "شناسه", "name": "شرح", "mode": "per_id"}') >= 2,
              "bot: request draft + warehouse return pass the material spec")
    else:
        print(f"  (request draft: {str(err)[:60]})")

    # assistant -------------------------------------------------------------------
    from services import assistant_brain as ab
    from services import assistant_tables as at

    for q in ("چی کم داریم", "موجودی استوپر", "مصرف ۳ ماه گذشته", "برای ۴ تاندیش اسلب چی لازمه", "حد تحمل نازل"):
        a = ab.answer(db, owner, q, use_llm=False, log=False, rephrase_answer=False)
        if not at.has_table(a.table):
            spy.take()
            continue
        files = at.write_files(a.table, out / "asst" / str(abs(hash(q))))
        collect(f"دستیار «{q}»", list(files.values()),
                expect_material=any(s.get("material") and s.get("rows") for s in a.table["sections"]))

    print(f"  reports checked: {len(reports)}")
    for name, files in list(reports.items())[:3]:
        print(f"    {name}: {[f.name for f in files]}")


def main() -> None:
    part_a()
    if HAVE_DB:
        part_b()
    else:
        print("== part B skipped (no DB) ==")
    if FAILS:
        print(f"FAILED {len(FAILS)}")
        sys.exit(1)
    print("OK smoke_material_columns")


if __name__ == "__main__":
    main()
