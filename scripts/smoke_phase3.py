"""Phase-3 smoke (1405-07-18 spec items 19–20). Temp dirs/DBs only — never live data."""
from __future__ import annotations

import _smoke_env  # noqa: F401,E402  — temp DB/reports/uploads before config import

import os
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
TMP = _smoke_env.SMOKE_TMP


def test_housekeeping() -> None:
    import config
    from services import housekeeping as hk

    assert str(config.REPORT_DIR).startswith(str(TMP)) and str(config.DATABASE_PATH).startswith(str(TMP))
    data = TMP / "hk_data"
    (data / "backups").mkdir(parents=True)
    hk.DATA_DIR, hk.BACKUP_DIR = data, data / "backups"
    now = time.time()
    for i in range(8):  # 8 legacy backups in data/, ages 0..70 days
        p = data / f"bot_backup_2026{i:02d}_x.db"
        p.write_bytes(b"x")
        os.utime(p, (now - i * 10 * 86400, now - i * 10 * 86400))
    (data / "backups" / "web_admin_credentials.txt").write_text("u")
    assert len(hk.consolidate_backups()) == 8 and not list(data.glob("bot_backup_*.db"))
    removed = {p.name for p in hk.prune_backups(now=now)}
    # newest 5 (0..40 d) kept; 50/60/70 d are >30 d and not newest-5 → removed
    assert removed == {"bot_backup_202605_x.db", "bot_backup_202606_x.db", "bot_backup_202607_x.db"}, removed
    assert (data / "backups" / "web_admin_credentials.txt").exists()
    rep = TMP / "hk_reports"
    (rep / "sub").mkdir(parents=True)
    old, new = rep / "sub" / "old.pdf", rep / "new.pdf"
    old.write_text("o"); new.write_text("n"); (rep / ".gitkeep").write_text("")
    os.utime(old, (now - 61 * 86400, now - 61 * 86400))
    os.utime(rep / ".gitkeep", (now - 99 * 86400, now - 99 * 86400))
    assert hk.purge_reports(root=rep, now=now) == 1 and new.exists() and (rep / ".gitkeep").exists()
    # rotating single log
    import logging
    from logging.handlers import RotatingFileHandler

    hk.DATA_DIR = TMP
    path = hk.setup_logging("smoke_bot")
    hs = logging.getLogger().handlers
    assert len(hs) == 1 and isinstance(hs[0], RotatingFileHandler)
    assert hs[0].maxBytes == 2 * 1024 * 1024 and hs[0].backupCount == 5 and path.parent == TMP
    # getUpdates read-timeouts are transient (WARNING, not ERROR)
    import main as bot_main
    from bot.bale_api import BaleAPIError

    assert bot_main._is_transient(BaleAPIError("getUpdates", "The read operation timed out"))
    assert not bot_main._is_transient(BaleAPIError("getUpdates", "Unauthorized"))
    logging.getLogger().handlers.clear()
    print("  19e housekeeping OK (backups keep 5/30d, reports 60d, rotating log, transient timeouts)")


def test_cache_and_audit() -> None:
    import sqlite3

    sys.path.insert(0, str(ROOT / "scripts"))
    import smoke_phase2 as p2
    import services.main_source as ms
    from bot.activity import log_activity
    from db.models import Database
    from services import frame_cache

    ms.UPLOAD_DIR = TMP / "uploads_ms"
    db = Database(TMP / "audit.db")
    db.upsert_user("31", role="owner", display_name="A")
    db.upsert_user("32", role="responsible_officer", display_name="B")
    p1 = ms.persist_primary_frame(db, p2._inv_frame(5.0), bale_user_id="31")
    n_rows = lambda: sqlite3.connect(db.path).execute("SELECT COUNT(*) FROM extracted_datasets").fetchone()[0]
    assert n_rows() == 1
    frame_cache.clear()
    ms.load_primary_frame(db); ms.load_primary_frame(db)
    assert frame_cache.stats == {"hits": 1, "misses": 1}, frame_cache.stats
    item = p2._inv_frame().iloc[0]["id"]
    # same user edits 3× → same clean_path → extract row updated, not inserted
    for q in (6, 7, 8):
        res = ms.upsert_row(db, item, {"quantity": q, "unit": "کیلو"}, bale_user_id="31")
    assert n_rows() == 1, n_rows()
    df = ms.load_primary_frame(db)  # cache key changed with mtime → fresh value
    assert float(df.loc[df["id"] == item, "quantity"].iloc[0]) == 8
    assert res["changes"] == [{"field": "quantity", "before": "7", "after": "8"}], res["changes"]
    first = ms.upsert_row(db, item, {"quantity": 9, "unit": "عدد"}, bale_user_id="31")["changes"]
    assert {c["field"] for c in first} == {"quantity", "unit"}
    txt = ms.changes_fa(first)
    assert "→" in txt and "8" in txt and "9" in txt, txt
    log_activity(db, db.get_user("31"), "edit_main_source_record", item_id=item, changes_fa=txt)
    msg = sqlite3.connect(db.path).execute("SELECT message_fa FROM user_activity ORDER BY id DESC").fetchone()[0]
    assert item in msg and "8 → 9" in msg, msg
    # another user's edit → own file → new extract row (and becomes factory-wide latest)
    ms.upsert_row(db, item, {"quantity": 1}, bale_user_id="32")
    assert n_rows() == 2
    # timestamped snapshots, newest N kept
    for q in range(15):
        ms.upsert_row(db, item, {"quantity": 100 + q}, bale_user_id="32")
    snaps = list((Path(ms.load_primary_frame.__globals__["resolve_primary_inventory_path"](db)).parent / "snapshots").glob("edit_snapshot_*.xlsx"))
    assert len(snaps) == ms.SNAPSHOT_KEEP, len(snaps)
    assert n_rows() == 2
    # add_row activity phrase is Persian (was raw key before)
    from bot.activity import format_action_phrase

    assert "افزود" in format_action_phrase("add_main_source_record", item_id="X", action_fa="رکورد جدید")
    print("  19a/19b OK (frame cache hits, extract row updated in place, edit audit before→after, snapshots keep 10)")


def test_errors_and_expiry() -> None:
    import logging

    sys.path.insert(0, str(ROOT / "scripts"))
    import smoke_phase1 as p1
    from bot import keyboards as kb
    from bot.handlers import BotApp
    from db.models import Database
    from services import user_errors

    db = Database(TMP / "err.db")
    db.upsert_user("41", role="owner", display_name="A")
    client = p1.FakeClient()
    app = BotApp(client, db)  # type: ignore[arg-type]
    catcher = p1.ErrCatcher()
    logging.getLogger().addHandler(catcher)

    def boom(_m):
        raise RuntimeError("secret internal detail /path/x.py")

    app.on_status = boom  # type: ignore[assignment]
    n = len(client.sent)
    app.handle_update({"update_id": 7, "message": {"from": {"id": 41, "first_name": "a"}, "chat": {"id": 41}, "text": "/status"}})
    txt = client.sent[-1][1]
    assert len(client.sent) == n + 1 and "کد پیگیری: E-" in txt and "secret" not in txt, txt
    code = txt.split("کد پیگیری: ")[1].split()[0]
    assert any(code in r and "secret internal detail" in r for r in catcher.records), catcher.records[-1:]
    # user-facing validation messages are kept verbatim (no code)
    m = user_errors.error_fa("ثبت نشد", ValueError("شناسه الزامی است."))
    assert "کد پیگیری" not in m and "شناسه الزامی است" in m
    logging.getLogger().removeHandler(catcher)

    # flow open → «restart» (new BotApp) → first typed value gets the expiry notice once
    def send(a, text):
        k = len(client.sent)
        a.handle_message({"from": {"id": 41, "first_name": "a"}, "chat": {"id": 41}, "text": text})
        return [t for _c, t, _m in client.sent[k:]]

    send(app, kb.BTN_BOT_SETTINGS); send(app, kb.BTN_USERS)
    send(app, kb.BTN_USERS_ADD)  # role picker = multi-step flow keyboard
    assert db.get_setting("flow_open:41"), "marker"
    time.sleep(0.01)
    app2 = BotApp(client, db)  # type: ignore[arg-type]
    out = send(app2, kb.BTN_ROLE_TECH)  # answer to the dead flow
    assert any("منقضی شد" in t for t in out), out
    assert not db.get_setting("flow_open:41")
    out = send(app2, kb.BTN_ROLE_TECH)
    assert not any("منقضی شد" in t for t in out), out
    # a fresh process where the user just presses a menu button → no notice
    send(app2, kb.BTN_BOT_SETTINGS); send(app2, kb.BTN_USERS); send(app2, kb.BTN_USERS_ADD)
    app3 = BotApp(client, db)  # type: ignore[arg-type]
    out = send(app3, kb.BTN_HOME)
    assert not any("منقضی شد" in t for t in out), out
    # web: unexpected error → 500 Persian page with tracking code, no detail
    from fastapi.testclient import TestClient

    import web.deps as deps
    from web.app import create_app

    wapp = create_app()
    wapp.dependency_overrides[deps.get_db] = lambda: db
    wapp.dependency_overrides[deps.current_user_optional] = lambda: db.get_user("41")

    @wapp.get("/__boom")
    async def _boom():
        raise RuntimeError("secret web detail")

    r = TestClient(wapp, raise_server_exceptions=False).get("/__boom")
    assert r.status_code == 500 and "کد پیگیری" in r.text and "secret" not in r.text, r.text[:300]
    print("  19d OK (tracking code bot+web, details only in log, expired-flow notice after restart)")


def check_xlsx(path: Path) -> list[str]:
    from openpyxl import load_workbook

    from excel.table_style import _first_table_row

    bad: list[str] = []
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if not ws.sheet_view.rightToLeft:
            bad.append(f"{path.name}/{ws.title}: not RTL")
        header = _first_table_row(ws)
        for row in ws.iter_rows():
            for c in row:
                if c.value in (None, ""):
                    continue
                if c.font.name != "Vazirmatn":
                    bad.append(f"{path.name}/{ws.title}!{c.coordinate} font {c.font.name}")
                if c.alignment.horizontal != "center":
                    bad.append(f"{path.name}/{ws.title}!{c.coordinate} align {c.alignment.horizontal}")
                if header and c.row >= header and c.border.left.style != "thin" and c.border.left.style is None:
                    bad.append(f"{path.name}/{ws.title}!{c.coordinate} no border")
        if header:
            hcells = [c for c in ws[header] if c.value not in (None, "")]
            if any(c.fill.fill_type is None for c in hcells):
                bad.append(f"{path.name}/{ws.title}: header row {header} without fill")
    return bad[:5]


def pdf_fonts(path: Path) -> list[str]:
    import subprocess

    out = subprocess.run(["pdffonts", str(path)], capture_output=True, text=True).stdout
    return [ln.split()[0].split("+")[-1] for ln in out.splitlines()[2:] if ln.strip()]


LETTERHEAD_FONTS: set[str] = set()


def check_pdf(path: Path) -> list[str]:
    names = pdf_fonts(path)
    report_fonts = [n for n in names if "Vazirmatn" in n]
    # fonts embedded in the user's own letterhead artwork are not ours to change
    others = [n for n in names if "Vazirmatn" not in n and n not in LETTERHEAD_FONTS]
    bad = []
    if not report_fonts:
        bad.append(f"{path.name}: no Vazirmatn ({names})")
    if any(x in n for n in others for x in ("Helvetica", "DejaVu", "Tahoma", "Times")):
        bad.append(f"{path.name}: other fonts {others}")
    return bad


def test_styling() -> None:
    import logging
    import shutil

    sys.path.insert(0, str(ROOT / "scripts"))
    import smoke_phase1 as p1
    from bot.handlers import BotApp
    from pdf import generator as gen

    gen._register_fonts()
    assert gen.FONT_NAME == "Vazirmatn" and gen.FONT_BOLD == "Vazirmatn-Bold", gen.FONT_NAME
    assert (ROOT / "fonts" / "Vazirmatn-Regular.ttf").is_file() and (ROOT / "fonts" / "Vazirmatn-Bold.ttf").is_file()

    tmp = TMP / "style"
    tmp.mkdir()
    db = p1._db_copy(tmp)
    owner = next(u for u in db.list_users() if u.get("role") == "owner" and u.get("active"))
    uid = str(owner["bale_user_id"])
    client = p1.FakeClient()
    app = BotApp(client, db)  # type: ignore[arg-type]
    catcher = p1.ErrCatcher()
    logging.getLogger().addHandler(catcher)
    lh = app._letterhead_path()
    if lh and Path(lh).suffix.lower() == ".pdf":
        LETTERHEAD_FONTS.update(pdf_fonts(Path(lh)))
    p1.walk_role(app, client, catcher, uid, "owner")
    # explicit generators not reached by the walk
    from excel.inbound import write_inbound_excel
    import pandas as pd

    write_inbound_excel(pd.DataFrame([{"کد کالا": "1234", "مقدار": 3}]), TMP / "reports" / "inb.xlsx")
    client.docs.append(("x", str(TMP / "reports" / "inb.xlsx")))
    from web.services import reports as wrep

    for fn in ("generate_monthly_summary_files", "generate_user_activity_files"):
        f = getattr(wrep, fn, None)
        if f:
            try:
                res = f(db, owner) if fn == "generate_monthly_summary_files" else f(db, owner, days=30)
                for x in res:
                    if isinstance(x, (str, Path)) and Path(x).suffix in {".pdf", ".xlsx"}:
                        client.docs.append(("x", str(x)))
            except TypeError:
                pass
    def _collect(obj):
        for v in (obj if isinstance(obj, (list, tuple)) else vars(obj).values() if hasattr(obj, "__dict__") else []):
            if isinstance(v, (str, Path)) and Path(v).suffix in {".pdf", ".xlsx"}:
                client.docs.append(("x", str(v)))
            elif isinstance(v, (list, tuple)):
                _collect(v)

    for name in ("generate_remaining_critical_files", "generate_surplus_files"):
        _collect(getattr(wrep, name)(db, owner, days=30))
    from services import critical_items_report as cir, period_consumption as pc
    from datetime import date, timedelta

    _collect(cir.generate_critical_items_files(db, owner))
    end = date.today()
    _collect(pc.generate_files(db, end - timedelta(days=30), end, range_label="۳۰ روز"))
    logging.getLogger().removeHandler(catcher)
    docs = sorted({Path(p) for _c, p in client.docs if Path(p).is_file()})
    xs = [d for d in docs if d.suffix == ".xlsx"]
    ps = [d for d in docs if d.suffix == ".pdf"]
    assert len(xs) >= 6 and len(ps) >= 6, (len(xs), len(ps))
    bad = [b for d in xs for b in check_xlsx(d)] + [b for d in ps for b in check_pdf(d)]
    assert not bad, bad[:15]
    keep = Path("/tmp/phase3_style_samples")
    shutil.rmtree(keep, ignore_errors=True)
    keep.mkdir()
    for d in docs:
        shutil.copy2(d, keep / d.name)
    print(f"  20 styling OK ({len(xs)} xlsx: RTL/Vazirmatn/center/border/header fill; {len(ps)} pdf: Vazirmatn) → {keep}")


def test_shared_core() -> None:
    import subprocess

    sys.path.insert(0, str(ROOT / "scripts"))
    import smoke_phase1 as p1
    from openpyxl import load_workbook

    from analytics.frames import load_frames
    from bot import keyboards as kb
    from bot.handlers import BotApp
    from services import comprehensive_report as cr

    tmp = TMP / "core"
    tmp.mkdir()
    db = p1._db_copy(tmp)
    owner = next(u for u in db.list_users() if u.get("role") == "owner" and u.get("active"))
    uid = str(owner["bale_user_id"])
    client = p1.FakeClient()
    app = BotApp(client, db)  # type: ignore[arg-type]
    # 19f: bot frames == analytics.frames.load_frames
    bf, _ = app._load_frames(owner, db.get_or_create_session(uid))
    wf = {k: v for k, v in load_frames(db, owner).items() if v is not None}
    assert set(bf) == set(wf) and all(bf[k].shape == wf[k].shape for k in wf), ({k: v.shape for k, v in bf.items()}, {k: v.shape for k, v in wf.items()})

    def send(text):
        app.handle_message({"from": {"id": int(uid), "first_name": "x"}, "chat": {"id": int(uid)}, "text": text})

    send(kb.BTN_ANALYTICS); send(kb.BTN_COMPREHENSIVE)
    nd = len(client.docs)
    send(kb.BTN_MY_CURRENT)
    if len(client.docs) == nd:  # section step
        send(next(iter(kb.SECTION_STEP_BUTTONS)))
    bot_docs = [Path(p) for _c, p in client.docs[nd:]]
    assert any(d.suffix == ".pdf" for d in bot_docs) and any(d.suffix == ".xlsx" for d in bot_docs), client.sent[-3:]
    bot_x = next(d for d in bot_docs if d.suffix == ".xlsx")
    bot_p = next(d for d in bot_docs if d.suffix == ".pdf")

    from fastapi.testclient import TestClient

    import web.deps as deps
    from web.app import create_app

    wapp = create_app()
    wapp.dependency_overrides[deps.get_db] = lambda: db
    who = {"u": owner}
    wapp.dependency_overrides[deps.current_user_optional] = lambda: who["u"]
    c = TestClient(wapp)
    html = c.get("/reports").text
    assert 'id="comprehensive"' in html and "/reports/comprehensive.xlsx" in html
    from bot.jalali import jalali_today

    t = jalali_today()
    q = f"from_year={t.year}&from_month={t.month}&to_year={t.year}&to_month={t.month}"
    r = c.get(f"/reports/comprehensive?{q}")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/pdf"), (r.status_code, r.text[:200])
    rx = c.get(f"/reports/comprehensive.xlsx?{q}")
    assert rx.status_code == 200 and rx.content[:2] == b"PK"
    wx = TMP / "web_comp.xlsx"
    wx.write_bytes(rx.content)
    wp = TMP / "web_comp.pdf"
    wp.write_bytes(r.content)
    assert load_workbook(bot_x).sheetnames == load_workbook(wx).sheetnames, (load_workbook(bot_x).sheetnames, load_workbook(wx).sheetnames)
    heads = lambda pdf: [ln for ln in subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout.splitlines() if ln.strip()][:3]
    assert heads(bot_p)[:1] == heads(wp)[:1], (heads(bot_p), heads(wp))
    assert not check_xlsx(wx) and not check_pdf(wp)
    # technician: denied on web (default permissions)
    db.upsert_user("777", role="technician", display_name="T")
    who["u"] = db.get_user("777")
    r = c.get(f"/reports/comprehensive?{q}", follow_redirects=False)
    assert r.status_code in (303, 403) or "دسترسی" in r.text, r.status_code
    print("  19f OK (bot frames = load_frames; bot+web comprehensive share core: same sheets/title; web route gated)")


def test_background() -> None:
    import threading

    sys.path.insert(0, str(ROOT / "scripts"))
    import smoke_phase1 as p1
    from bot import keyboards as kb
    from bot.handlers import BotApp
    from services import comprehensive_report as cr

    tmp = TMP / "bg"
    tmp.mkdir()
    db = p1._db_copy(tmp)
    owner = next(u for u in db.list_users() if u.get("role") == "owner" and u.get("active"))
    uid = str(owner["bale_user_id"])
    db.upsert_user("555", role="manager", display_name="M")

    class TSClient(p1.FakeClient):
        def __init__(self):
            super().__init__()
            self.log: list[tuple[str, str, str]] = []  # (thread, chat, text/doc)
            self._l = threading.Lock()

        def send_message(self, chat_id, text, reply_markup=None, **k):
            with self._l:
                self.log.append((threading.current_thread().name, str(chat_id), text))
                return super().send_message(chat_id, text, reply_markup, **k)

        def send_document(self, chat_id, path, caption=None, **k):
            with self._l:
                self.log.append((threading.current_thread().name, str(chat_id), "DOC " + str(path)))
                return super().send_document(chat_id, path, caption, **k)

    client = TSClient()
    app = BotApp(client, db)  # type: ignore[arg-type]
    app.bg.enable(workers=2)
    gate = threading.Event()
    real = cr.generate_files

    def slow(*a, **k):
        gate.wait(10)
        return real(*a, **k)

    cr.generate_files = slow
    upd = iter(range(1000, 2000))

    def send(who, text):
        app.handle_update({"update_id": next(upd), "message": {"from": {"id": int(who), "first_name": "x"}, "chat": {"id": int(who)}, "text": text}})

    try:
        send(uid, kb.BTN_ANALYTICS); send(uid, kb.BTN_COMPREHENSIVE)
        t0 = time.monotonic()
        send(uid, kb.BTN_MY_CURRENT)
        if not app.bg.busy(uid):  # section step first
            send(uid, next(iter(kb.SECTION_STEP_BUTTONS)))
        assert time.monotonic() - t0 < 2 and app.bg.busy(uid), "polling thread must not block"
        n_owner = len([x for x in client.log if x[1] == uid])
        send(uid, kb.BTN_HOME)  # queued while the job runs
        time.sleep(0.2)
        assert len([x for x in client.log if x[1] == uid]) == n_owner + 1 or any("در حال ساخت" in x[2] for x in client.log[-3:]), client.log[-3:]
        before_other = len(client.log)
        send("555", "/start")  # another user is served immediately
        assert any(x[1] == "555" for x in client.log[before_other:]), "other user blocked"
        gate.set()
        for _ in range(300):
            if not app.bg.busy(uid):
                break
            time.sleep(0.1)
        assert not app.bg.busy(uid)
        mine = [x for x in client.log if x[1] == uid]
        docs = [i for i, x in enumerate(mine) if x[2].startswith("DOC ")]
        home = [i for i, x in enumerate(mine) if "منوی اصلی" in x[2] or "خانه" in x[2]]
        assert docs and any("در حال ساخت" in x[2] for x in mine), mine[-6:]
        assert all(x[0].startswith("heavy") for x in mine if x[2].startswith("DOC ")), "docs not from worker"
        assert home and max(home) > max(docs), ("queued HOME must run after the job", mine[-5:])
        # background failure → tracking code to the user, app keeps working
        cr.generate_files = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("bg secret"))
        send(uid, kb.BTN_ANALYTICS); send(uid, kb.BTN_COMPREHENSIVE); send(uid, kb.BTN_MY_CURRENT)
        if not app.bg.busy(uid):
            send(uid, next(iter(kb.SECTION_STEP_BUTTONS)))
        for _ in range(100):
            if not app.bg.busy(uid):
                break
            time.sleep(0.05)
        last = [x[2] for x in client.log if x[1] == uid][-3:]
        assert any("کد پیگیری" in t for t in last) and not any("bg secret" in t for t in last), last
    finally:
        cr.generate_files = real
        app.bg.shutdown()
    print("  19c OK (heavy report in worker thread, polling not blocked, per-user queue in order, other users served, errors → tracking code)")


def main() -> int:
    print("smoke phase3…")
    test_housekeeping()
    test_cache_and_audit()
    test_errors_and_expiry()
    test_styling()
    test_shared_core()
    test_background()
    print("SMOKE_PHASE3_OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
