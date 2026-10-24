"""Smoke v1.2.0: 📥 ورود نرخ‌های مصرف + rate-loss guard + «نرخ ندارد» counts (no live data written)."""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd

from bot import keyboards as kb
from bot.activity import ACTION_PHRASES, format_action_phrase
from services import main_source as ms
from services import rates_import as ri


def _src() -> pd.DataFrame:
    return ms.ensure_inventory_columns(pd.DataFrame({
        "category_code": ["1201", "1201", "1302", "1400"],
        "id": ["120100000001", "120100000002", "130200000001", "140000000001"],
        "product_name": ["a", "b", "c", "d"],
        "quantity": [1, 2, 3, 4],
        "slab_patching": [None, 5, None, None],
        "casting_floor": [None, None, 2, None],
    }))


def main() -> None:
    src = _src()
    # export → re-import round trip (Tahoma/dark header handled by shared table_style)
    with tempfile.TemporaryDirectory() as td:
        out = ri.export_rates_xlsx(src, Path(td) / "r.xlsx")
        back = ri.read_rates_file(out)
        assert set(ri.RATE_COLS) <= set(back.columns) and len(back) == 2
        from openpyxl import load_workbook

        c = load_workbook(out).active["A1"]
        assert c.font.name == "Tahoma" and c.alignment.horizontal == "center"
        incoming = pd.DataFrame({
            "category_code": ["1201", "1201", "1302", "9999"],
            "id": ["120100000001", "120100000002", "", ""],
            "slab_patching": [3, 7, None, 1],
            "casting_floor": [None, None, 2, None],
            "billet_patching": [None, None, 4, None],
        })
        plan = ri.build_plan(src, incoming)
        assert len(plan.fills) == 2, plan.fills          # id1 slab_patching + code 1302 billet_patching
        assert len(plan.conflicts) == 1 and plan.conflicts[0]["old"] == 5
        assert plan.same == 1 and len(plan.unmatched) == 1
        assert "مغایرت" in ri.summary_fa(plan)
        rev = ri.review_xlsx(plan, Path(td) / "rev.xlsx")
        assert load_workbook(rev).sheetnames == ["پر می‌شوند", "مغایرت", "بی‌تطابق"]
        assert ri.Plan.from_dict(plan.to_dict()).to_dict() == plan.to_dict()
    # rates presence message (N تاندیش)
    msg = ri.presence_message_fa(src, "slab", "without")
    assert "از ۴ قلم منبع اصلی، ۱ قلم نرخ پچینگ اسلب دارند" in msg, msg
    assert "۰ قلم نرخ نوسازی/پچینگ بیلت" in ri.presence_message_fa(src, "billet", "with")
    # guard: casting_floor counts; partial wipe detected
    only_cf = src.copy(); only_cf["slab_patching"] = None
    blank = only_cf.copy(); blank["casting_floor"] = None
    assert ms.rates_wiped(only_cf, blank) and not ms.rates_wiped(None, blank)
    loss = ms.rate_losses(src, only_cf)
    assert len(loss) == 1 and loss[0]["col"] == "slab_patching"
    assert "تأیید پاک شدن نرخ‌ها" in ms.rate_loss_note_fa(loss)
    assert ms.no_rates_warning_fa(blank) and not ms.no_rates_warning_fa(src)
    added = [{"id": "140000000001", "category_code": "1400", "product_name": "d"}]
    assert len(ms.added_without_rates(src, added)) == 1
    # UI wiring
    owner = {"role": "owner", "bale_user_id": "1"}
    flat = [b for r in kb.main_source_file_menu(owner)["keyboard"] for b in r]
    flat = [b.get("text", b) if isinstance(b, dict) else b for b in flat]
    assert kb.BTN_RATES_IMPORT in flat
    assert kb.rates_hint_inline(owner)["inline_keyboard"][0][0]["callback_data"] == "ri|start"
    assert "rates_import" in ACTION_PHRASES and "{" not in format_action_phrase("rates_import", filled=3, overwritten=0)
    from bot.handlers import BotApp

    for name in ("on_rates_import_start", "on_rates_import_document", "_on_rates_callback", "on_rate_loss_ok"):
        assert hasattr(BotApp, name), name
    from web.routers.settings import router

    paths = {getattr(r, "path", "") for r in router.routes}
    assert {"/settings/main-source/rates/review", "/settings/main-source/rates/apply"} <= paths
    _web_flow()
    print("OK smoke_rates_import")


def _web_flow() -> None:
    """Review + cancel only (never applies to the live منبع اصلی)."""
    from fastapi.testclient import TestClient

    import web.deps as deps
    from db.models import Database
    from web.app import create_app
    from web.auth_web import set_credential

    deps.reset_db_singleton()
    db = Database()
    db.upsert_user("8801", "owner", display_name="مالک آزمون")
    set_credential(db, bale_user_id="8801", username="owner8801", password="secret88")
    deps._db = db
    app = create_app()
    app.dependency_overrides[deps.get_db] = lambda: db
    c = TestClient(app)
    assert c.post("/login", data={"username": "owner8801", "password": "secret88"},
                  follow_redirects=False).status_code == 303
    page = c.get("/settings/main-source").text
    assert "ورود نرخ‌های مصرف" in page and "allow_rate_loss" in page
    frame = ms.load_primary_frame(db, bale_user_id="8801")
    if frame is None or ri.rows_with_rates(frame).empty:
        return
    with tempfile.TemporaryDirectory() as td:
        f = ri.export_rates_xlsx(frame, Path(td) / "x.xlsx")
        r = c.post("/settings/main-source/rates/review",
                   files={"file": ("x.xlsx", f.read_bytes(), "application/octet-stream")})
        assert r.status_code == 200 and "بررسی فایل نرخ‌های مصرف" in r.text, r.text[-800:]
        assert c.get("/settings/main-source/rates/review.xlsx").status_code == 200
        r = c.post("/settings/main-source/rates/apply", data={"mode": "cancel"})
        assert "لغو شد" in r.text


if __name__ == "__main__":
    main()
