#!/usr/bin/env python3
"""Smoke (temp DB only): a stock upload that lacks main-source IDs keeps their stock
unchanged and raises a Persian alarm (کد ۴ رقمی | آیدی | شرح | موجودی قبلی | واحد) to the
uploader + active owner/manager; no duplicate per file; >15 → Tahoma/centered Excel."""
from __future__ import annotations

import _smoke_env  # noqa: F401,E402

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def main() -> int:
    import pandas as pd
    from openpyxl import load_workbook

    from _smoke_env import SMOKE_TMP
    from db.models import Database
    from services import main_source as ms
    from services import unit_mismatch as um

    old = pd.DataFrame([
        {"category_code": "1603", "id": "378721562302U", "product_name": "نسوز خوزستان KRCAST LAN 65",
         "quantity": 136, "unit": "No"},
        {"category_code": "1473", "id": "378471411121U", "product_name": "SOLWOOL", "quantity": 11500, "unit": "No"},
    ])
    new = pd.DataFrame([{"category_code": "1473", "id": "378471411121U", "product_name": "SOLWOOL", "quantity": 11000}])
    merged, info = ms.merge_stock_update(old, new)
    k = merged.set_index("id")
    assert float(k.at["378721562302U", "quantity"]) == 136  # kept, not zeroed
    assert float(k.at["378471411121U", "quantity"]) == 11000
    assert [m["آیدی"] for m in info["missing"]] == ["378721562302U"], info

    db = Database()
    for uid, role in (("7001", "owner"), ("7002", "manager"), ("7003", "technician")):
        db.upsert_user(uid, role, display_name=role)
    raw = SMOKE_TMP / "stock.xlsx"
    new.to_excel(raw, index=False)
    a = um.check_missing_stock(db, raw_path=raw, items=info["missing"], file_label="به‌روزرسانی موجودی انبار")
    assert a, "alarm expected"
    t = a["text"]
    for must in ("🚨 هشدار: ۱ قلم منبع اصلی در فایل «به‌روزرسانی موجودی انبار» نبود",
                 "موجودی این اقلام تغییر نکرد", "(کد ۴ رقمی | آیدی | شرح | موجودی قبلی | واحد):",
                 "• 1603 | 378721562302U | نسوز خوزستان KRCAST LAN 65 | ۱۳۶ | No"):
        assert must in t, (must, t)
    assert um.check_missing_stock(db, raw_path=raw, items=info["missing"], file_label="x") is None  # no duplicate

    class C:
        def __init__(self):
            self.to = []

        def send_message(self, cid, *a, **k):
            self.to.append(cid)

        def send_document(self, *a, **k):
            pass

    c = C()
    um.send_alert(c, db, a, uploader="7003")
    assert sorted(c.to) == ["7001", "7002", "7003"], c.to

    many = [{"کد ۴ رقمی": "1203", "آیدی": f"I{i:02d}", "شرح": "ش", "موجودی قبلی": i, "واحد": "Kg"} for i in range(16)]
    raw2 = SMOKE_TMP / "stock2.xlsx"
    pd.DataFrame([{"x": 1}]).to_excel(raw2, index=False)
    b = um.check_missing_stock(db, raw_path=raw2, items=many, file_label="به‌روزرسانی موجودی انبار")
    assert b and b["xlsx"] and "فهرست کامل در فایل اکسل پیوست" in b["text"]
    ws = load_workbook(b["xlsx"]).active
    assert ws.max_row == 17 and ws["D2"].font.name == "Tahoma" and ws["D2"].alignment.horizontal == "center"
    print("smoke_stock_missing_alarm OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
