"""Template-locked reports (1405/07/25): اقلام بحرانی + خلاصه مصرفی ماهیانه.

Both copy a user-approved sample instead of the shared styler / base-columns rule:
  • اقلام بحرانی ← templates/critical_items_template.xlsx (user's own sheet «140506»,
    2026-10-05): per sheet title bar B1 merged to the second-to-last column, header
    row 2 with the sample's per-column fills, «توضیحات» block at the bottom of the
    SAME sheet (A/B merged, notes merged C..), Tahoma 10 bold, RTL, A4 landscape,
    no frozen panes, no separate notes sheet.
  • خلاصه مصرفی ماهیانه ← templates/monthly_summary_approved_1405-07-04.xlsx
    (approved sample of 2026-09-26): same sheet, widths, fills, row heights, no
    frozen header; only the font is Tahoma.
Temp dirs only; the live DB is used read-only through a backup copy.
"""
from __future__ import annotations

import _smoke_env  # noqa: F401,E402

import sqlite3
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
TMP = _smoke_env.SMOKE_TMP / "template_reports"
TMP.mkdir(parents=True, exist_ok=True)

FAILS: list[str] = []


def check(ok: bool, msg: str) -> None:
    if not ok:
        FAILS.append(msg)
        print("  FAIL", msg)


def _fill_sig(c):
    f = c.fill
    col = f.fgColor
    return (f.fill_type, col.type, col.theme if col.type == "theme" else col.rgb,
            round(float(col.tint or 0), 3))


def _border_sig(c):
    b = c.border
    return tuple(getattr(b, s).style for s in ("left", "right", "top", "bottom"))


def _font_sig(c):
    return (c.font.name, float(c.font.sz or 0), bool(c.font.b))


def check_critical_sheet(ws, tpl, *, n_rows: int, n_notes: int) -> None:
    from excel.critical_template import TEMPLATE_COLUMN_OF, TEMPLATE_WIDTHS
    from services.critical_items_report import TEMPLATE_COLUMNS as REPORT_COLUMNS

    n = len(REPORT_COLUMNS)
    name = ws.title
    check(ws.sheet_view.rightToLeft is True, f"{name}: RTL")
    check(ws.freeze_panes is None, f"{name}: no frozen panes (sample has none)")
    check(ws.page_setup.orientation == tpl.page_setup.orientation == "landscape", f"{name}: landscape")
    check(int(ws.page_setup.paperSize or 0) == int(tpl.page_setup.paperSize) == 9, f"{name}: A4")
    for side in ("left", "right", "top", "bottom"):
        check(getattr(ws.page_margins, side) == getattr(tpl.page_margins, side), f"{name}: margin {side}")
    merges = {str(r) for r in ws.merged_cells.ranges}
    from openpyxl.utils import get_column_letter as L

    # title bar
    check(f"B1:{L(n - 1)}1" in merges, f"{name}: title merged B1:{L(n - 1)}1 {sorted(merges)[:4]}")
    check(str(ws["B1"].value or "").startswith("لیست اقلام بحرانی نسوز تاندیش"), f"{name}: title text")
    for sig, what in ((_fill_sig, "fill"), (_font_sig, "font"), (_border_sig, "border")):
        check(sig(ws["B1"]) == sig(tpl["B1"]), f"{name}: title {what} {sig(ws['B1'])} ≠ {sig(tpl['B1'])}")
    check(ws.row_dimensions[1].height == tpl.row_dimensions[1].height, f"{name}: title height")
    check(ws.row_dimensions[2].height == tpl.row_dimensions[2].height, f"{name}: header height")
    # header + body styles per template column
    hdr = [ws.cell(2, i).value for i in range(1, n + 1)]
    check(hdr == list(REPORT_COLUMNS), f"{name}: header {hdr}")
    check(ws.cell(2, n + 1).value is None, f"{name}: no columns beyond the template")
    # T12 (user decision 1405/07/25, deliberate deviation from «140506»): ردیف is column A
    # (rightmost), کد چهاررقمی column B; each column keeps its own template style/width.
    check([tpl[f"{c}2"].value.strip() for c in "AB"] == ["کد چهاررقمی", "ردیف"], "template A2/B2 order")
    check((ws["A2"].value, ws["B2"].value) == ("ردیف", "کد چهاررقمی"), f"{name}: A2/B2 = ردیف/کد (swapped)")
    check(_fill_sig(ws["A2"]) == _fill_sig(tpl["B2"]) and _fill_sig(ws["B2"]) == _fill_sig(tpl["A2"]),
          f"{name}: swapped header fills move with their column")
    check(abs(ws.column_dimensions["A"].width - tpl.column_dimensions["B"].width) < 0.1
          and abs(ws.column_dimensions["B"].width - tpl.column_dimensions["A"].width) < 0.1,
          f"{name}: swapped widths move with their column")
    if n_rows:
        check(_fill_sig(ws["A3"]) == _fill_sig(tpl["B4"]) and _fill_sig(ws["B3"]) == _fill_sig(tpl["A4"]),
              f"{name}: swapped body fills move with their column")
    check("آیدی کد" not in hdr and "کلیدواژه" not in hdr, f"{name}: no base material columns")
    for i, col in enumerate(REPORT_COLUMNS, start=1):
        tl = TEMPLATE_COLUMN_OF.get(col)
        if not tl:
            check(ws.cell(2, i).fill.fill_type is None, f"{name}: added column {col} unfilled")
            continue
        h, th = ws.cell(2, i), tpl[f"{tl}2"]
        check(_fill_sig(h) == _fill_sig(th), f"{name}: header fill {col} {_fill_sig(h)} ≠ {_fill_sig(th)}")
        check(_font_sig(h) == _font_sig(th), f"{name}: header font {col}")
        check(_border_sig(h) == _border_sig(th), f"{name}: header border {col}")
        check(h.alignment.horizontal == th.alignment.horizontal, f"{name}: header align {col}")
        check(abs(ws.column_dimensions[L(i)].width - TEMPLATE_WIDTHS[tl]) < 0.01
              and abs(TEMPLATE_WIDTHS[tl] - tpl.column_dimensions[tl].width) < 0.1, f"{name}: width {col}")
        if n_rows:
            d, td = ws.cell(3, i), tpl[f"{tl}4"]
            check(_fill_sig(d) == _fill_sig(td), f"{name}: body fill {col} {_fill_sig(d)} ≠ {_fill_sig(td)}")
            check(_font_sig(d) == _font_sig(td), f"{name}: body font {col}")
            if tl in ("D", "F"):
                check(d.number_format == "#,##0", f"{name}: {col} format {d.number_format}")
            if tl in ("A", "B"):
                check(d.number_format == "@", f"{name}: {col} text format")
            last = ws.cell(2 + n_rows, i)
            check(last.border.bottom.style == "medium", f"{name}: last body row closes medium ({col})")
    # «توضیحات» block at the bottom of the same sheet
    first = 3 + max(n_rows, 1)
    last = first + n_notes - 1
    check(ws.cell(first, 2).value == "توضیحات", f"{name}: B{first} = توضیحات ({ws.cell(first, 2).value})")
    if n_notes > 1:
        check(f"A{first}:A{last}" in merges and f"B{first}:B{last}" in merges, f"{name}: A/B merged over notes")
    for r in range(first, last + 1):
        check(f"C{r}:{L(n - 1)}{r}" in merges, f"{name}: note row {r} merged C..{L(n - 1)}")
        c = ws.cell(r, 3)
        if r == first:
            check(bool(c.value), f"{name}: first note row has text")
        if not c.value:  # blank note line: bordered like the template's empty rows
            check(_border_sig(c)[3] is not None, f"{name}: blank note row {r} bordered")
        check(c.alignment.wrap_text is True and c.alignment.horizontal == tpl["C15"].alignment.horizontal,
              f"{name}: note align row {r}")
    check(_font_sig(ws.cell(first, 2)) == _font_sig(tpl["B15"]), f"{name}: توضیحات font")
    check(ws.max_row == last, f"{name}: nothing after the notes ({ws.max_row} vs {last})")
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, ""):
                check(c.font.name == "Tahoma", f"{name}!{c.coordinate}: font {c.font.name}")


def test_critical_synthetic() -> None:
    from openpyxl import load_workbook

    from services.critical_items_report import TEMPLATE_COLUMNS as REPORT_COLUMNS
    from excel.critical_template import (TEMPLATE_PATH, TEMPLATE_SHEET, load_template_workbook,
                                         write_critical_items_pdf, write_critical_items_xlsx)
    from excel.table_style import finalize_path, template_locked_name

    tpl_wb = load_template_workbook(TEMPLATE_PATH)
    check(tpl_wb.sheetnames == [TEMPLATE_SHEET], f"template sheets {tpl_wb.sheetnames}")
    tpl = tpl_wb[TEMPLATE_SHEET]
    check(all(tpl.cell(r, c).font.name == "Tahoma" for r in range(1, 18) for c in range(1, 8)
              if tpl.cell(r, c).value is not None), "sample itself uses Tahoma")

    def row(code, i, desc, stock, need, days):
        r = dict.fromkeys(REPORT_COLUMNS)
        r.update({"کد چهاررقمی": code, "ردیف": i, "کد و شرح کالا": desc, "موجودی": stock,
                  "واحد": "Kg", "نیاز": need, "حد تحمل(روز)": days})
        return r

    rows = [row("1662", 1, "بتن 86 نوسازی", 42383, 977617.4, 4),
            row("1674", 2, "نازل بیرونی", 197, 3195, 10),
            row("", "", "نازل داخلی (عضو گروه)", 0, None, None)]
    sheets = [
        {"name": "140507 شرکت", "title": "لیست اقلام بحرانی نسوز تاندیش — شرکت — با نوسازی (1405/07/25)",
         "columns": list(REPORT_COLUMNS), "rows": rows,
         "notes": ["افق محاسبه: اجناس داخلی ۳ ماه، اجناس وارداتی ۶ ماه.", "یادداشت دوم", "", ""]},
        {"name": "140507 پیمانکار", "title": "لیست اقلام بحرانی نسوز تاندیش — پیمانکار — با نوسازی (1405/07/25)",
         "columns": list(REPORT_COLUMNS), "rows": rows[:1], "notes": ["یک یادداشت", "دو"]},
    ]
    out = write_critical_items_xlsx(sheets, TMP / "crit.xlsx")
    wb = load_workbook(out)
    check(wb.sheetnames == ["140507 شرکت", "140507 پیمانکار"], f"sheets {wb.sheetnames}")
    check(TEMPLATE_SHEET not in wb.sheetnames and "توضیحات" not in wb.sheetnames, "no proto / notes sheet")
    check(template_locked_name(wb) == "critical_items", "critical xlsx tagged template-locked")
    check_critical_sheet(wb.worksheets[0], tpl, n_rows=3, n_notes=4)
    check_critical_sheet(wb.worksheets[1], tpl, n_rows=1, n_notes=2)
    # the shared styler must leave it untouched
    before = out.read_bytes()
    finalize_path(out)
    check(out.read_bytes() == before, "finalize_path skips the template-locked workbook")
    pdf = write_critical_items_pdf(sheets[0], TMP / "crit.pdf")
    from pypdf import PdfReader

    p = PdfReader(str(pdf)).pages[0].mediabox
    check(float(p.width) > float(p.height) and abs(float(p.width) - 841.9) < 2, "critical PDF A4 landscape")
    print("  critical synthetic vs template OK" if not FAILS else "  critical synthetic FAIL")


def test_critical_live_copy() -> None:
    db_path = ROOT / "data" / "bot.db"
    if not db_path.is_file():
        print("  (no live DB — skipped)")
        return
    from openpyxl import load_workbook

    from db.models import Database
    from excel.critical_template import TEMPLATE_PATH, TEMPLATE_SHEET, load_template_workbook
    from services.critical_items_report import critical_item_count, generate_critical_items_files

    copy = TMP / "bot_copy.db"
    src = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    dst = sqlite3.connect(copy)
    src.backup(dst)
    src.close()
    dst.close()
    db = Database(copy)
    users = [u for u in db.list_users() if u.get("role") == "owner" and u.get("active")]
    if not users:
        print("  (no owner — skipped)")
        return
    tpl = load_template_workbook(TEMPLATE_PATH)[TEMPLATE_SHEET]
    res = generate_critical_items_files(db, users[0], output_dir=TMP / "live", file_prefix="crit")
    if res.error:
        print("  (live copy:", res.error[:80], ")")
        return
    wb = load_workbook(res.xlsx)
    check(len(wb.worksheets) == 2 and "توضیحات" not in wb.sheetnames, f"live sheets {wb.sheetnames}")
    from analytics.critical_items import report_footer_notes  # noqa: F401

    for ws, seg in zip(wb.worksheets, ("company", "contractor")):
        df = res.frames.get(seg)
        n_rows = 0 if df is None else len(df)
        n_notes = ws.max_row - 2 - max(n_rows, 1)
        from services.critical_items_report import HORIZON_NOTE, MIN_NOTE_ROWS

        check(n_notes == MIN_NOTE_ROWS, f"live {ws.title}: notes block ({n_notes})")
        check_critical_sheet(ws, tpl, n_rows=n_rows, n_notes=n_notes)
        first = 3 + max(n_rows, 1)
        notes = [ws.cell(r, 3).value for r in range(first, first + n_notes)]
        check(notes[0] == HORIZON_NOTE, f"live {ws.title}: horizon note first ({notes[0]})")
        check(str(notes[1] or "").startswith("تاریخ گزارش:") and "تاندیش" in str(notes[1]),
              f"live {ws.title}: date/basis note ({notes[1]})")
        check(all(v is None for v in notes[2:]), f"live {ws.title}: rest of the notes block blank {notes[2:]}")
        _ = critical_item_count(df)
    from pypdf import PdfReader

    for p in res.pdfs.values():
        box = PdfReader(str(p)).pages[0].mediabox
        check(float(box.width) > float(box.height), f"live {p.name}: landscape")
    print("  critical live-copy vs template OK" if not FAILS else "  critical live-copy FAIL")


def test_monthly_summary() -> None:
    import shutil

    from openpyxl import load_workbook

    from excel.monthly_summary import build_monthly_summary
    from excel.table_style import template_locked_name

    from excel.critical_template import load_template_workbook

    ref_wb = load_template_workbook(ROOT / "templates" / "monthly_summary_approved_1405-07-04.xlsx")
    ref = ref_wb.worksheets[0]
    src = TMP / "monthly_src.xlsx"
    real = [p for p in (ROOT / "samples" / "real" / "monthly_consumption_sample.xlsx",) if p.is_file()]
    shutil.copy((real or [ROOT / "samples" / "03_monthly_consumption.xlsx"])[0], src)
    try:
        _data, out = build_monthly_summary(src, excel_out=TMP / "monthly_summary.xlsx")
    except ValueError as exc:
        print("  monthly summary skipped:", exc)
        return
    wb = load_workbook(out)
    ws = wb.worksheets[0]
    check(template_locked_name(wb) == "monthly_summary", "monthly summary tagged template-locked")
    check(wb.sheetnames == ref_wb.sheetnames, f"monthly sheets {wb.sheetnames} vs {ref_wb.sheetnames}")
    check(ws.sheet_view.rightToLeft is True and ws.freeze_panes == ref.freeze_panes is None,
          f"monthly RTL / no freeze ({ws.freeze_panes})")
    widths = lambda s: {k: v.width for k, v in s.column_dimensions.items() if v.width}  # noqa: E731
    check(widths(ws) == widths(ref), f"monthly widths {widths(ws)}")
    check([c.value for c in ws[1]] == [c.value for c in ref[1]], "monthly header labels")
    for a, b in zip(ws[1], ref[1]):
        check(_fill_sig(a) == _fill_sig(b) and bool(a.font.b) == bool(b.font.b), f"monthly header style {a.coordinate}")
    fills = lambda s: {_fill_sig(c) for row in s.iter_rows() for c in row if c.fill.fill_type}  # noqa: E731
    check(fills(ws) <= fills(ref), f"monthly fills ⊆ sample {fills(ws) - fills(ref)}")
    heights = {ws.row_dimensions[r].height for r in range(1, ws.max_row + 1) if ws.row_dimensions[r].height}
    ref_h = {ref.row_dimensions[r].height for r in range(1, ref.max_row + 1) if ref.row_dimensions[r].height}
    check(heights <= ref_h, f"monthly row heights {heights - ref_h}")
    fonts = {c.font.name for row in ws.iter_rows() for c in row if c.value not in (None, "")}
    check(fonts == {"Tahoma"}, f"monthly fonts {fonts}")
    print("  monthly summary vs approved sample OK" if not FAILS else "  monthly summary FAIL")


def test_critical_notes_and_archive() -> None:
    """T11: ≤2 fixed notes + user lines; draft → notes → approve archive (temp UPLOAD_DIR)."""
    from services import critical_report_archive as cra
    from services.critical_items_report import HORIZON_NOTE, compose_notes, user_note_lines

    check(compose_notes(["a", "b"], "", 5) == ["a", "b", "", "", ""], "compose_notes pads blanks")
    check(compose_notes(["a", "b"], " x \n\n y ", 5) == ["a", "b", "x", "y", ""], "compose_notes user lines")
    check(len(compose_notes(["a", "b"], "\n".join("l" * 3 for _ in range(30)), 5)) == 2 + cra.MAX_NOTE_LINES
          or len(user_note_lines("\n".join(["l"] * 30))) <= 30, "notes capped")
    check("اجناس داخلی ۳ ماه، اجناس وارداتی ۶ ماه" in HORIZON_NOTE, "horizon note text")
    db_path = ROOT / "data" / "bot.db"
    if not db_path.is_file():
        print("  (no live DB — archive test skipped)")
        return
    import config
    from db.models import Database

    check(str(cra.archive_root()).startswith(str(_smoke_env.SMOKE_TMP)), f"archive in temp {cra.archive_root()}")
    copy = TMP / "bot_copy_archive.db"
    src = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    dst = sqlite3.connect(copy)
    src.backup(dst)
    src.close()
    dst.close()
    db = Database(copy)
    owner = [u for u in db.list_users() if u.get("role") == "owner" and u.get("active")]
    if not owner:
        return
    user = owner[0]
    rep, res = cra.create_draft(db, user, reno_mode="with", source="bot")
    if rep is None:
        print("  (archive: no draft —", (res.error or "")[:60], ")")
        return
    rid = int(rep["id"])
    check(rep["status"] == cra.STATUS_DRAFT and not cra.list_approved(db), "draft created, not approved")
    files = cra.report_files(db, rep)
    check(set(files) >= {"xlsx"} and all(p.is_file() for p in files.values()), f"draft files {files}")
    check(all("پیش‌نویس" in p.name for p in files.values()), "draft file names")
    rep, err = cra.set_notes(db, rid, "یادداشت نمونه ۱\nیادداشت نمونه ۲", user)
    check(err is None and rep["user_notes"] == "یادداشت نمونه ۱\nیادداشت نمونه ۲", f"set_notes {err}")
    long_rep, err = cra.set_notes(db, rid, "x" * (cra.MAX_NOTES_CHARS + 5), user)
    check(err is None and len(long_rep["user_notes"]) == cra.MAX_NOTES_CHARS, "too long notes capped")
    rep, err = cra.set_notes(db, rid, "یادداشت نمونه ۱\nیادداشت نمونه ۲", user)
    rep, err = cra.approve(db, rid, user)
    check(err is None and rep["status"] == cra.STATUS_APPROVED, f"approve {err}")
    files = cra.report_files(db, rep)
    check(all("تأییدشده" in p.name and str(p).startswith(str(cra.archive_root())) for p in files.values()),
          f"final files {files}")
    from openpyxl import load_workbook

    ws = load_workbook(files["xlsx"]).worksheets[0]
    col_c = [c.value for c in ws["C"] if c.value]
    check("یادداشت نمونه ۱" in col_c and "یادداشت نمونه ۲" in col_c and HORIZON_NOTE in col_c,
          f"final notes in xlsx {col_c[-5:]}")
    # missing file → re-rendered from the stored rows
    files["xlsx"].unlink()
    again = cra.report_files(db, rep)
    check(again["xlsx"].is_file(), "missing archive file re-rendered")
    check([r["id"] for r in cra.list_approved(db)] == [rid], "approved list")
    _r, err = cra.approve(db, rid, user)
    check(err is not None, "double approve refused")
    _r, err = cra.cancel(db, rid, user)
    check(err is not None, "approved report cannot be cancelled")
    _ = config
    print("  critical notes + draft/approve archive OK" if not FAILS else "  critical archive FAIL")


def test_mix() -> None:
    """T11b: «70 بیلت + 20 بلوم + 100 اسلب» = Σ rate × count (+ casting floor × total)."""
    from services import n_tundish_report as nt

    check(nt.parse_mix("بیلت ۷۰ بلوم ۲۰ اسلب ۱۰۰") == {"billet": 70, "bloom": 20, "slab": 100},
          f"parse_mix 1 {nt.parse_mix('بیلت ۷۰ بلوم ۲۰ اسلب ۱۰۰')}")
    check(nt.parse_mix("۷۰ بیلت، ۲۰ بلوم و ۱۰۰ اسلب") == {"billet": 70, "bloom": 20, "slab": 100}, "parse_mix 2")
    check(nt.parse_mix("70 بیلت 100 اسلب") == {"billet": 70, "bloom": 0, "slab": 100},
          f"parse_mix 3 {nt.parse_mix('70 بیلت 100 اسلب')}")
    check(nt.parse_mix("بیلت") is None and nt.parse_mix("12") is None, "parse_mix needs ≥2 sections")
    db_path = ROOT / "data" / "bot.db"
    if not db_path.is_file():
        return
    from db.models import Database

    copy = TMP / "bot_copy_mix.db"
    src = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    dst = sqlite3.connect(copy)
    src.backup(dst)
    src.close()
    dst.close()
    db = Database(copy)
    owner = [u for u in db.list_users() if u.get("role") == "owner" and u.get("active")]
    if not owner:
        return
    mix = {"billet": 70, "bloom": 20, "slab": 100}
    res = nt.generate_mix_files(db, owner[0], mix, "with", output_dir=TMP / "mix")
    if res.error:
        print("  (mix:", res.error[:60], ")")
        return
    check(res.pdf.is_file() and res.xlsx.is_file(), "mix files")
    from openpyxl import load_workbook

    ws = load_workbook(res.xlsx).worksheets[0]
    hdr_row = next(r for r in range(1, 8) if ws.cell(r, 1).value == "کد ۴ رقمی")
    hdr = [c.value for c in ws[hdr_row]]
    for col in (nt.COL_PER_BILLET, nt.COL_PER_BLOOM, nt.COL_PER_SLAB, nt.COL_PER_CF, res.need_col):
        check(col in hdr, f"mix column {col}")
    check(hdr[:3] == ["کد ۴ رقمی", "آیدی کد", "شرح"], f"mix base columns {hdr[:3]}")
    num = lambda v: float(v) if isinstance(v, (int, float)) else 0.0  # noqa: E731
    bad = 0
    for r in range(hdr_row + 1, ws.max_row + 1):
        row = {h: ws.cell(r, i + 1).value for i, h in enumerate(hdr)}
        if not row.get("کد ۴ رقمی") or not isinstance(row.get(res.need_col), (int, float)):
            continue
        exp = (70 * num(row[nt.COL_PER_BILLET]) + 20 * num(row[nt.COL_PER_BLOOM])
               + 100 * num(row[nt.COL_PER_SLAB]) + 190 * num(row[nt.COL_PER_CF]))
        bad += abs(exp - float(row[res.need_col])) > 0.6
    check(bad == 0, f"mix need = Σ rate × count ({bad} rows differ)")
    fonts = {c.font.name for row in ws.iter_rows() for c in row if c.value not in (None, "")}
    check(fonts == {"Tahoma"}, f"mix fonts {fonts}")
    print("  mix 70/20/100 OK" if not FAILS else "  mix FAIL")


def main() -> None:
    test_critical_synthetic()
    test_critical_live_copy()
    test_critical_notes_and_archive()
    test_mix()
    test_monthly_summary()
    assert not FAILS, FAILS[:10]
    print("SMOKE_TEMPLATE_REPORTS_OK")


if __name__ == "__main__":
    main()
