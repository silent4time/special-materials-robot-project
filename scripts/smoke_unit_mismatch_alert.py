#!/usr/bin/env python3
"""Smoke (temp DB only): upload-time unit check vs منبع اصلی (bot + shared web path).

SH/No and KG/No mismatches + items missing from the main source are listed in a Persian
warning to the uploader and every active owner/manager; the same file+item+unit pair is
not re-alerted; >15 items → Tahoma/centered Excel; no data changed."""
from __future__ import annotations

import _smoke_env  # noqa: F401,E402

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def main() -> int:
    import pandas as pd
    from openpyxl import load_workbook

    from _smoke_env import SMOKE_TMP
    from db.models import Database
    from services import unit_mismatch as um

    maps = ({"378471411121U": "No", "378700009162X": "No", "378721562302U": "No", "A1": "Kg"},
            {"1473": "No", "1603": "No"})
    frame = pd.DataFrame([
        {"category_code": "1473", "id": "378471411121U", "material_name": "378471411121U - SOLWOOL", "unit": "SH"},
        {"category_code": "1473", "id": "378471411121U", "material_name": "378471411121U - SOLWOOL", "unit": "SH"},
        {"category_code": "1473", "id": "378700009162X", "material_name": "378700009162X - فلت", "unit": "SH"},
        {"category_code": "1603", "id": "378721562302U", "material_name": "378721562302U - KRCAST", "unit": "KG"},
        {"category_code": "1700", "id": "A1", "material_name": "A1 - ok", "unit": "KG"},  # Kg≈KG → no alert
        {"category_code": "9999", "id": "ZZ1", "material_name": "ZZ1 - ناشناس", "unit": "SET"},  # missing
    ])
    rows, missing = um.find_mismatches(frame, maps)
    got = {(r["آیدی"], r["واحد منبع اصلی"], r["واحد فایل"], r["تعداد ردیف"]) for r in rows}
    assert got == {("378471411121U", "No", "SH", 2), ("378700009162X", "No", "SH", 1),
                   ("378721562302U", "No", "KG", 1)}, got
    assert [r["آیدی"] for r in missing] == ["ZZ1"], missing
    text = um.alert_text_fa(rows, missing, file_label="مصرف ماهیانه مواد")
    for must in ("⚠️ مغایرت واحد", "کد ۴ رقمی | آیدی | شرح | واحد منبع اصلی | واحد فایل | تعداد ردیف",
                 "1473 | 378471411121U | SOLWOOL | No | SH | ۲ ردیف", "واحد «منبع اصلی»", "در «منبع اصلی» نیست"):
        assert must in text, (must, text)

    db = Database()
    db.upsert_user("6001", "owner", display_name="مالک")
    db.upsert_user("6002", "manager", display_name="مدیر")
    db.upsert_user("6003", "technician", display_name="تکنسین")
    raw = SMOKE_TMP / "m.xlsx"
    frame.to_excel(raw, index=False)
    a1 = um.check_upload(db, raw_path=raw, clean_path=raw, file_type="monthly_consumption", maps=maps,
                         file_label="مصرف ماهیانه مواد")
    assert a1 and len(a1["rows"]) == 3
    a2 = um.check_upload(db, raw_path=raw, clean_path=raw, file_type="monthly_consumption", maps=maps,
                         file_label="مصرف ماهیانه مواد")
    assert a2 is None, "same file + item + unit must not re-alert"

    class C:
        def __init__(self):
            self.to = []

        def send_message(self, cid, text, *a, **k):
            self.to.append(cid)

        def send_document(self, *a, **k):
            pass

    c = C()
    um.send_alert(c, db, a1, uploader="6003")
    assert sorted(c.to) == ["6001", "6002", "6003"], c.to  # uploader + owner + manager

    big = [{"کد ۴ رقمی": "1", "آیدی": f"X{i}", "شرح": "ش", "واحد منبع اصلی": "No", "واحد فایل": "SH",
            "تعداد ردیف": 1} for i in range(16)]
    x = um.write_excel(big, [], SMOKE_TMP / "u.xlsx")
    ws = load_workbook(x).active
    assert ws.max_row == 17 and ws["D2"].font.name == "Tahoma" and ws["D2"].alignment.horizontal == "center"
    assert ws.sheet_view.rightToLeft
    assert list(frame["unit"])[:1] == ["SH"]  # input untouched
    print("smoke_unit_mismatch_alert OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
