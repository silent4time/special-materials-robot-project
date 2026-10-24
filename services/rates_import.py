"""📥 ورود نرخ‌های مصرف — import the 7 rate columns into منبع اصلی from an Excel file.

Shared by the bot and the web panel. Accepts the export format
(کد دسته / شناسه / شرح / واحد + 7 rate columns) or a full منبع اصلی file.
Match = شناسه مواد first, then the 4-digit کد دسته (applies to every row of that code
that was not matched by id). Nothing is written until ``apply_plan`` is called after the
user confirms; a DB backup is taken first.
"""
from __future__ import annotations

import logging
import math
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd

from excel.processor import _normalize_columns
from services import main_source as ms

logger = logging.getLogger(__name__)

RATE_COLS: tuple[str, ...] = (
    "slab_renovation", "slab_patching", "bloom_renovation", "bloom_patching",
    "billet_renovation", "billet_patching", "casting_floor",
)
EXPORT_COLS: list[str] = ["category_code", "id", "product_name", "unit", *RATE_COLS]


def _num(v: Any) -> float | None:
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip().translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩٫", "01234567890123456789."))
        s = s.replace(",", "").replace("،", "")
        if not s or s in {"-", "nan", "None"}:
            return None
        try:
            v = float(s)
        except ValueError:
            return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if math.isnan(f):
        return None
    return f


def _has_rate(v: Any) -> bool:
    f = _num(v)
    return f is not None and f != 0


def _code(v: Any) -> str:
    s = ms._norm_id(v)
    return s[:4] if len(s) >= 4 and s[:4].isdigit() else s


def rows_with_rates(frame: pd.DataFrame) -> pd.DataFrame:
    """Rows having at least one non-zero rate (7 columns)."""
    if frame is None or frame.empty:
        return pd.DataFrame(columns=EXPORT_COLS)
    work = ms.ensure_inventory_columns(frame.copy())
    mask = pd.Series(False, index=work.index)
    for c in RATE_COLS:
        if c in work.columns:
            mask |= work[c].map(_has_rate)
    return work.loc[mask]


def export_rates_xlsx(frame: pd.DataFrame, out: Path | str) -> Path:
    """Tahoma / centered / dark header (shared table_style) export of rows with rates."""
    from excel.simple_report import export_dataframe_xlsx

    sub = rows_with_rates(frame)
    if sub.empty:
        raise ValueError("هیچ ردیفی با نرخ مصرف در منبع اصلی نیست.")
    return export_dataframe_xlsx(sub, out, columns=EXPORT_COLS, header_map=ms.FIELD_LABELS_FA,
                                 sheet_name="نرخ مصرف")


@dataclass
class Plan:
    fills: list[dict] = field(default_factory=list)      # {id, code, name, col, new}
    conflicts: list[dict] = field(default_factory=list)  # {id, code, name, col, old, new}
    unmatched: list[dict] = field(default_factory=list)  # {id, code, name}
    same: int = 0
    file_rows: int = 0
    error: str = ""

    @property
    def ok(self) -> bool:
        return not self.error

    def to_dict(self) -> dict:
        return {"fills": self.fills, "conflicts": self.conflicts, "unmatched": self.unmatched,
                "same": self.same, "file_rows": self.file_rows, "error": self.error}

    @classmethod
    def from_dict(cls, d: dict) -> "Plan":
        return cls(fills=list(d.get("fills") or []), conflicts=list(d.get("conflicts") or []),
                   unmatched=list(d.get("unmatched") or []), same=int(d.get("same") or 0),
                   file_rows=int(d.get("file_rows") or 0), error=str(d.get("error") or ""))


def read_rates_file(path: Path | str) -> pd.DataFrame:
    """Read the first sheet whose header row maps to id/code + at least one rate column."""
    xl = pd.ExcelFile(path, engine="openpyxl")
    for sheet in xl.sheet_names:
        raw = pd.read_excel(xl, sheet_name=sheet, header=None, dtype=object)
        for hdr in range(min(10, len(raw))):
            df = pd.read_excel(xl, sheet_name=sheet, header=hdr, dtype=object)
            df = _normalize_columns(df)
            cols = set(map(str, df.columns))
            if ({"id", "category_code"} & cols) and (set(RATE_COLS) & cols):
                return df
    raise ValueError("ستون شناسه/کد دسته و حداقل یک ستون نرخ مصرف در فایل پیدا نشد.")


def build_plan(source: pd.DataFrame | None, incoming: pd.DataFrame) -> Plan:
    plan = Plan()
    if source is None or source.empty:
        plan.error = "منبع اصلی هنوز بارگذاری نشده است."
        return plan
    src = ms.ensure_inventory_columns(source.copy())
    rate_cols = [c for c in RATE_COLS if c in incoming.columns]
    if not rate_cols:
        plan.error = "هیچ ستون نرخ مصرفی در فایل نیست."
        return plan
    src_ids = src["id"].map(ms._norm_id)
    src_codes = src["category_code"].map(_code)
    by_id: dict[str, list[int]] = {}
    by_code: dict[str, list[int]] = {}
    for idx, (i, c) in enumerate(zip(src_ids, src_codes)):
        if i:
            by_id.setdefault(i, []).append(idx)
        if c:
            by_code.setdefault(c, []).append(idx)
    id_matched: set[int] = set()
    rows = []
    for _, r in incoming.iterrows():
        vals = {c: _num(r.get(c)) for c in rate_cols}
        vals = {c: v for c, v in vals.items() if v is not None and v != 0}
        if not vals:
            continue
        rid = ms._norm_id(r.get("id")) if "id" in incoming.columns else ""
        rcode = _code(r.get("category_code")) if "category_code" in incoming.columns else ""
        # v1.3.0: material IDs do NOT start with the 4-digit code (e.g. 378121641302R) —
        # an unknown ID without a code stays unmatched instead of hitting a wrong category
        name = str(r.get("product_name") or "").strip()
        if name.lower() == "nan":
            name = ""
        rows.append((rid, rcode, name, vals))
    plan.file_rows = len(rows)
    targets: dict[int, dict[str, float]] = {}
    pending_code = []
    for rid, rcode, name, vals in rows:
        if rid and rid in by_id:
            for idx in by_id[rid]:
                id_matched.add(idx)
                targets.setdefault(idx, {}).update(vals)
        else:
            pending_code.append((rid, rcode, name, vals))
    for rid, rcode, name, vals in pending_code:
        idxs = [i for i in by_code.get(rcode, []) if i not in id_matched]
        if not idxs:
            plan.unmatched.append({"id": rid, "code": rcode, "name": name})
            continue
        for idx in idxs:
            t = targets.setdefault(idx, {})
            for c, v in vals.items():
                t.setdefault(c, v)
    for idx in sorted(targets):
        row = src.iloc[idx]
        base = {"id": ms._norm_id(row.get("id")), "code": _code(row.get("category_code")),
                "name": str(row.get("product_name") or "").strip(), "row": int(idx)}
        for c, new in targets[idx].items():
            old = _num(row.get(c))
            if old is None or old == 0:
                plan.fills.append({**base, "col": c, "new": new})
            elif abs(old - new) > 1e-9:
                plan.conflicts.append({**base, "col": c, "old": old, "new": new})
            else:
                plan.same += 1
    return plan


def plan_from_file(db: Any, user: dict, path: Path | str) -> Plan:
    try:
        incoming = read_rates_file(path)
    except Exception as exc:  # noqa: BLE001
        return Plan(error=f"فایل خوانده نشد: {exc}")
    source = ms.load_primary_frame(db, bale_user_id=user["bale_user_id"])
    return build_plan(source, incoming)


def summary_fa(plan: Plan) -> str:
    from services.units import fmt_num

    if plan.error:
        return "⚠️ " + plan.error
    lines = [
        "📥 بررسی فایل نرخ‌های مصرف",
        f"ردیف‌های دارای نرخ در فایل: {plan.file_rows}",
        f"✅ خانه‌های خالی که پر می‌شوند: {len(plan.fills)}",
        f"⚠️ مغایرت (مقدار فعلی متفاوت است): {len(plan.conflicts)} — پیش‌فرض: مقدار فعلی حفظ می‌شود",
        f"= بدون تغییر (مقدار یکسان): {plan.same}",
        f"❓ ردیف‌های بی‌تطابق (شناسه/کد در منبع اصلی نیست): {len(plan.unmatched)}",
    ]
    for c in plan.conflicts[:8]:
        lines.append(f"  • {c['id'] or c['code']} — {ms.FIELD_LABELS_FA.get(c['col'], c['col'])}: "
                     f"فعلی {fmt_num(c['old'])} ← فایل {fmt_num(c['new'])}")
    if len(plan.conflicts) > 8:
        lines.append(f"  … و {len(plan.conflicts) - 8} مغایرت دیگر (فایل اکسل بررسی)")
    lines.append("تا تأیید نکنید چیزی ذخیره نمی‌شود.")
    return "\n".join(lines)


def review_xlsx(plan: Plan, out: Path | str) -> Path:
    """Review workbook: پرشونده / مغایرت / بی‌تطابق sheets (shared table style)."""
    from openpyxl import Workbook

    from excel.contrast import fix_workbook_contrast
    from excel.table_style import style_workbook_path

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    lab = ms.FIELD_LABELS_FA
    wb = Workbook()
    ws = wb.active
    ws.title = "پر می‌شوند"
    ws.append(["کد دسته", "شناسه", "شرح", "ستون", "مقدار جدید"])
    for f in plan.fills:
        ws.append([f["code"], f["id"], f["name"], lab.get(f["col"], f["col"]), f["new"]])
    ws = wb.create_sheet("مغایرت")
    ws.append(["کد دسته", "شناسه", "شرح", "ستون", "مقدار فعلی", "مقدار فایل"])
    for f in plan.conflicts:
        ws.append([f["code"], f["id"], f["name"], lab.get(f["col"], f["col"]), f["old"], f["new"]])
    ws = wb.create_sheet("بی‌تطابق")
    ws.append(["کد دسته", "شناسه", "شرح"])
    for f in plan.unmatched:
        ws.append([f["code"], f["id"], f["name"]])
    wb.save(out)
    for name in ("پر می‌شوند", "مغایرت", "بی‌تطابق"):
        try:
            style_workbook_path(out, sheet_name=name)
        except Exception:  # noqa: BLE001
            logger.warning("style review sheet %s failed", name)
    try:
        from openpyxl import load_workbook

        wb2 = load_workbook(out)
        fix_workbook_contrast(wb2)
        wb2.save(out)
    except Exception:  # noqa: BLE001
        pass
    return out


def apply_plan(db: Any, user: dict, plan: Plan, *, overwrite: bool = False) -> dict:
    """Backup DB, then write fills (+ conflicts when overwrite) into منبع اصلی."""
    from services.housekeeping import backup_db

    if plan.error:
        return {"ok": False, "error": plan.error}
    changes = list(plan.fills) + (list(plan.conflicts) if overwrite else [])
    if not changes:
        return {"ok": True, "filled": 0, "overwritten": 0, "backup": ""}
    source = ms.load_primary_frame(db, bale_user_id=user["bale_user_id"])
    if source is None or source.empty:
        return {"ok": False, "error": "منبع اصلی پیدا نشد."}
    try:
        bpath = str(backup_db(Path(db.path), "pre_rates_import"))
    except Exception as exc:  # noqa: BLE001
        logger.exception("backup before rates import failed")
        return {"ok": False, "error": f"پشتیبان‌گیری پیش از ذخیره ناموفق بود؛ چیزی ذخیره نشد ({exc})."}
    work = ms.ensure_inventory_columns(source.copy())
    ids = work["id"].map(ms._norm_id).tolist()
    for c in RATE_COLS:
        if c in work.columns:
            work[c] = work[c].astype(object)
    done = 0
    for ch in changes:
        idx = int(ch.get("row", -1))
        # source may have changed since review → re-locate by id
        if not (0 <= idx < len(work)) or ids[idx] != ch["id"]:
            try:
                idx = ids.index(ch["id"])
            except ValueError:
                continue
        work.iat[idx, work.columns.get_loc(ch["col"])] = ch["new"]
        done += 1
    ms.persist_primary_frame(db, work, bale_user_id=user["bale_user_id"])
    return {"ok": True, "filled": sum(1 for _ in plan.fills),
            "overwritten": len(plan.conflicts) if overwrite else 0, "written": done, "backup": bpath,
            "ts": time.time()}


SECTION_FA = {"slab": "اسلب", "bloom": "بلوم", "billet": "بیلت"}


def _section_cols(section: str | None, reno_mode: str | None) -> list[str]:
    if not section:
        return list(RATE_COLS)
    cols = [f"{section}_patching"]
    if str(reno_mode or "with") != "without":
        cols.insert(0, f"{section}_renovation")
    return cols


def rates_presence(frame: pd.DataFrame | None, section: str | None = None,
                   reno_mode: str | None = None) -> tuple[int, int]:
    """(rows in منبع اصلی, rows having a rate) — optionally one section/mode."""
    if frame is None or frame.empty:
        return 0, 0
    cols = [c for c in _section_cols(section, reno_mode) if c in frame.columns]
    if not cols:
        return int(len(frame)), 0
    mask = pd.Series(False, index=frame.index)
    for c in cols:
        mask |= frame[c].map(_has_rate)
    return int(len(frame)), int(mask.sum())


def presence_message_fa(frame: pd.DataFrame | None, section: str | None = None,
                        reno_mode: str | None = None) -> str:
    """e.g. «از ۱۴۹ قلم منبع اصلی، ۰ قلم نرخ پچینگ اسلب دارند.» + what to do."""
    from services.units import fmt_num

    total, have = rates_presence(frame, section, reno_mode)
    if section:
        kind = ("نوسازی/پچینگ " if str(reno_mode or "with") != "without" else "پچینگ ") + SECTION_FA.get(section, section)
    else:
        kind = "مصرف"
    fa = lambda n: fmt_num(n).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))  # noqa: E731
    return (f"از {fa(total)} قلم منبع اصلی، {fa(have)} قلم نرخ {kind} دارند.\n"
            "نرخ‌ها را با «📥 ورود نرخ‌های مصرف» از فایل وارد کنید یا در «📦 منبع اصلی» ویرایش کنید.")
