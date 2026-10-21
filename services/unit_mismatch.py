"""Upload-time unit check: every input file's units vs منبع اصلی (by ID, then 4-digit code).

Read-only — nothing is changed; reports already use the main-source unit
(``services.unit_source``). After an upload (bot + web) the uploader and owner/manager
get a Persian warning listing mismatches (کد ۴ رقمی، آیدی، شرح، واحد منبع اصلی، واحد
فایل، تعداد ردیف) and items missing from the main source. The same item+unit pair is
never re-alerted for the same file (file content hash).
"""
from __future__ import annotations

import hashlib
import json
import logging
from pathlib import Path
from typing import Any

import pandas as pd

from services.unit_source import _norm_code, _norm_id, item_id_of

logger = logging.getLogger(__name__)

EXCEL_THRESHOLD = 15
SEEN_KEY = "unit_alert_seen:"
TEXT_ROWS = 15


def _fa(n: Any) -> str:
    return str(n).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))


def _same(a: str, b: str) -> bool:
    return a.strip().casefold() == b.strip().casefold()


def find_mismatches(frame: pd.DataFrame | None, maps: tuple[dict, dict]) -> tuple[list[dict], list[dict]]:
    """(mismatch rows, missing-from-main-source rows) aggregated per item+file unit."""
    if frame is None or frame.empty or "unit" not in frame.columns:
        return [], []
    by_id, by_code = maps
    n = len(frame)
    col = lambda c: frame[c] if c in frame.columns else [None] * n  # noqa: E731
    desc_col = "description" if "description" in frame.columns else ("material_name" if "material_name" in frame.columns else None)
    agg: dict[tuple, dict] = {}
    missing: dict[str, dict] = {}
    for iid_raw, code, unit, name, desc in zip(col("id"), col("category_code"), frame["unit"], col("material_name"),
                                              col(desc_col) if desc_col else [None] * n):
        fu = str(unit or "").strip()
        if not fu or fu.lower() == "nan":
            continue
        iid = item_id_of(iid_raw, name)
        c = _norm_code(code)
        d = str(desc or name or "").strip()
        if " - " in d and iid and d.upper().startswith(iid):
            d = d.split(" - ", 1)[1].strip()
        if iid and iid in by_id:
            mu = by_id[iid]
        elif c and c in by_code:
            mu = by_code[c]
        else:
            if iid or c:
                k = iid or c
                m = missing.setdefault(k, {"کد ۴ رقمی": c, "آیدی": iid, "شرح": d, "واحد فایل": fu, "تعداد ردیف": 0})
                m["تعداد ردیف"] += 1
            continue
        if _same(mu, fu):
            continue
        k = (iid or c, fu)
        a = agg.setdefault(k, {"کد ۴ رقمی": c, "آیدی": iid, "شرح": d, "واحد منبع اصلی": mu,
                               "واحد فایل": fu, "تعداد ردیف": 0})
        a["تعداد ردیف"] += 1
    rows = sorted(agg.values(), key=lambda r: (str(r["کد ۴ رقمی"]), str(r["آیدی"])))
    return rows, sorted(missing.values(), key=lambda r: (str(r["کد ۴ رقمی"]), str(r["آیدی"])))


def _file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:32]


def _dedupe(db: Any, raw_path: Path, rows: list[dict], missing: list[dict]) -> tuple[list[dict], list[dict]]:
    """Drop item+unit pairs already alerted for this exact file; remember the new ones."""
    try:
        key = SEEN_KEY + _file_hash(raw_path)
        seen = set(json.loads(db.get_setting(key) or "[]"))
    except Exception:  # noqa: BLE001
        return rows, missing
    pk = lambda r, tag: f"{tag}|{r.get('آیدی') or r.get('کد ۴ رقمی')}|{r.get('واحد فایل')}"  # noqa: E731
    new_rows = [r for r in rows if pk(r, "m") not in seen]
    new_missing = [r for r in missing if pk(r, "x") not in seen]
    seen |= {pk(r, "m") for r in new_rows} | {pk(r, "x") for r in new_missing}
    try:
        db.set_setting(key, json.dumps(sorted(seen), ensure_ascii=False))
    except Exception:  # noqa: BLE001
        logger.warning("unit alert dedupe write failed")
    return new_rows, new_missing


def alert_text_fa(rows: list[dict], missing: list[dict], *, file_label: str, actor: str = "",
                  has_excel: bool = False) -> str:
    lines = [f"⚠️ مغایرت واحد در فایل «{file_label}»" + (f" — {actor}" if actor else "")]
    if rows:
        lines.append(f"واحد {_fa(len(rows))} قلم با «منبع اصلی» یکسان نیست "
                     "(کد ۴ رقمی | آیدی | شرح | واحد منبع اصلی | واحد فایل | تعداد ردیف):")
        for r in rows[:TEXT_ROWS]:
            lines.append(f"• {r['کد ۴ رقمی'] or '—'} | {r['آیدی'] or '—'} | {str(r['شرح'])[:45]} | "
                         f"{r['واحد منبع اصلی']} | {r['واحد فایل']} | {_fa(r['تعداد ردیف'])} ردیف")
        if len(rows) > TEXT_ROWS:
            lines.append(f"… و {_fa(len(rows) - TEXT_ROWS)} قلم دیگر" + (" (فهرست کامل در فایل اکسل پیوست)" if has_excel else ""))
        lines.append("گزارش‌ها همیشه واحد «منبع اصلی» را نشان می‌دهند؛ داده‌ها تغییری نکرد. "
                     "اگر واحد منبع اصلی اشتباه است، آن را در منبع اصلی اصلاح کنید.")
    if missing:
        lines.append("")
        lines.append(f"ℹ️ {_fa(len(missing))} قلم این فایل در «منبع اصلی» نیست؛ برای آن‌ها واحد خود فایل استفاده می‌شود:")
        for r in missing[:TEXT_ROWS]:
            lines.append(f"• {r['کد ۴ رقمی'] or '—'} | {r['آیدی'] or '—'} | {str(r['شرح'])[:45]} | "
                         f"واحد فایل {r['واحد فایل']} | {_fa(r['تعداد ردیف'])} ردیف")
        if len(missing) > TEXT_ROWS:
            lines.append(f"… و {_fa(len(missing) - TEXT_ROWS)} قلم دیگر")
    return "\n".join(lines)


def write_excel(rows: list[dict], missing: list[dict], out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheets = [("مغایرت واحد", ["ردیف", "کد ۴ رقمی", "آیدی", "شرح", "واحد منبع اصلی", "واحد فایل", "تعداد ردیف"], rows)]
    if missing:
        sheets.append(("خارج از منبع اصلی", ["ردیف", "کد ۴ رقمی", "آیدی", "شرح", "واحد فایل", "تعداد ردیف"], missing))
    for i, (title, cols, data) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.sheet_view.rightToLeft = True
        ws.append(cols)
        for n, r in enumerate(data, 1):
            ws.append([n] + [r.get(c) for c in cols[1:]])
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Tahoma", size=10, bold=c.row == 1)
                c.alignment = center
                c.border = border
                if c.row == 1:
                    c.fill = PatternFill("solid", fgColor="D9E1F2")
        for letter, w in zip("ABCDEFG", (6, 11, 18, 50, 15, 12, 11)):
            ws.column_dimensions[letter].width = w
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _frame_for(raw_path: Path, clean_path: Path | None, file_type: str) -> pd.DataFrame | None:
    """The uploaded file's own rows (units as written in the file)."""
    from excel.processor import load_excel

    order = [raw_path, clean_path] if file_type == "product_inventory" else [clean_path, raw_path]
    for p in order:
        if not p or not Path(p).is_file():
            continue
        try:
            f = load_excel(p, strict_tundish=False) if p == raw_path else pd.read_excel(p, dtype=str)
        except Exception:  # noqa: BLE001
            continue
        if f is not None and "unit" in f.columns and ({"id", "category_code"} & set(f.columns)):
            return f
    return None


def check_upload(db: Any, *, raw_path: Path, clean_path: Path | None, file_type: str, maps: tuple[dict, dict],
                 file_label: str, actor: str = "") -> dict | None:
    """Alert dict {text, xlsx, rows, missing} for NEW mismatches of this file, or None."""
    if not (maps[0] or maps[1]):
        return None
    frame = _frame_for(Path(raw_path), Path(clean_path) if clean_path else None, file_type)
    rows, missing = find_mismatches(frame, maps)
    if file_type == "product_inventory":
        missing = []  # a stock upload already reports new / unknown IDs itself
    if not rows and not missing:
        return None
    rows, missing = _dedupe(db, Path(raw_path), rows, missing)
    if not rows and not missing:
        return None
    xlsx = None
    if len(rows) + len(missing) > EXCEL_THRESHOLD:
        try:
            from config import REPORT_DIR
            from services.file_names import fa_stem, unique_dir

            xlsx = write_excel(rows, missing, unique_dir(REPORT_DIR, "unit_mismatch") / f"{fa_stem('مغایرت واحد', file_label)}.xlsx")
        except Exception:  # noqa: BLE001
            logger.exception("unit mismatch excel failed")
    text = alert_text_fa(rows, missing, file_label=file_label, actor=actor, has_excel=bool(xlsx))
    return {"text": text, "xlsx": str(xlsx) if xlsx else None, "rows": rows, "missing": missing}


def managers_chat_ids(db: Any, exclude: str | None = None) -> list[str]:
    out = []
    for u in db.list_users(active_only=True):
        if u.get("role") in {"owner", "manager"} and str(u.get("bale_user_id")) != str(exclude or ""):
            out.append(str(u["bale_user_id"]))
    return out


def send_alert(client: Any, db: Any, alert: dict, *, uploader: str | None, to_uploader: bool = True) -> int:
    """Uploader (optional) + every active owner/manager; Excel attached when present."""
    ids = ([str(uploader)] if to_uploader and uploader else []) + managers_chat_ids(db, exclude=uploader)
    sent = 0
    for cid in ids:
        try:
            client.send_message(cid, alert["text"])
            if alert.get("xlsx"):
                try:
                    client.send_document(cid, alert["xlsx"], caption="فهرست کامل مغایرت واحد با منبع اصلی")
                except Exception:  # noqa: BLE001
                    logger.warning("unit alert excel send failed")
            sent += 1
        except Exception:  # noqa: BLE001
            logger.warning("unit alert send failed")
    return sent


# ---- stock upload: main-source IDs missing from the new file (stock kept unchanged) ----
def _qty_fa(v: Any) -> str:
    try:
        f = float(v)
        return _fa(int(f) if f == int(f) else round(f, 2))
    except (TypeError, ValueError):
        return "—"


def missing_stock_text_fa(items: list[dict], *, file_label: str, actor: str = "", has_excel: bool = False) -> str:
    lines = [
        f"🚨 هشدار: {_fa(len(items))} قلم منبع اصلی در فایل «{file_label}» نبود" + (f" — {actor}" if actor else ""),
        "موجودی این اقلام تغییر نکرد (همان موجودی قبلی ماند) و همچنان در گزارش‌ها حساب می‌شود. "
        "اگر قلمی واقعاً حذف یا تمام شده، آن را در منبع اصلی اصلاح کنید.",
        "(کد ۴ رقمی | آیدی | شرح | موجودی قبلی | واحد):",
    ]
    for r in items[:TEXT_ROWS]:
        lines.append(f"• {r.get('کد ۴ رقمی') or '—'} | {r.get('آیدی') or '—'} | {str(r.get('شرح') or '')[:45]} | "
                     f"{_qty_fa(r.get('موجودی قبلی'))} | {r.get('واحد') or '—'}")
    if len(items) > TEXT_ROWS:
        lines.append(f"… و {_fa(len(items) - TEXT_ROWS)} قلم دیگر" + (" (فهرست کامل در فایل اکسل پیوست)" if has_excel else ""))
    return "\n".join(lines)


def write_missing_excel(items: list[dict], out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "اقلام نبوده در فایل"
    ws.sheet_view.rightToLeft = True
    cols = ["ردیف", "کد ۴ رقمی", "آیدی", "شرح", "موجودی قبلی", "واحد"]
    ws.append(cols)
    for n, r in enumerate(items, 1):
        ws.append([n] + [r.get(c) for c in cols[1:]])
    thin = Side(style="thin")
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Tahoma", size=10, bold=c.row == 1)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if c.row == 1:
                c.fill = PatternFill("solid", fgColor="F8CBAD")
    for letter, w in zip("ABCDEF", (6, 11, 18, 50, 14, 10)):
        ws.column_dimensions[letter].width = w
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def check_missing_stock(db: Any, *, raw_path: Path, items: list[dict], file_label: str, actor: str = "") -> dict | None:
    """Alert dict for main-source IDs absent from a stock upload (new ones only per file)."""
    if not items:
        return None
    try:
        key = SEEN_KEY + _file_hash(Path(raw_path))
        seen = set(json.loads(db.get_setting(key) or "[]"))
        new = [r for r in items if f"gone|{r.get('آیدی')}" not in seen]
        db.set_setting(key, json.dumps(sorted(seen | {f"gone|{r.get('آیدی')}" for r in new}), ensure_ascii=False))
    except Exception:  # noqa: BLE001
        new = items
    if not new:
        return None
    new = sorted(new, key=lambda r: (str(r.get("کد ۴ رقمی")), str(r.get("آیدی"))))
    xlsx = None
    if len(new) > EXCEL_THRESHOLD:
        try:
            from config import REPORT_DIR
            from services.file_names import fa_stem, unique_dir

            xlsx = write_missing_excel(new, unique_dir(REPORT_DIR, "stock_missing") / f"{fa_stem('اقلام نبوده در فایل موجودی')}.xlsx")
        except Exception:  # noqa: BLE001
            logger.exception("missing-stock excel failed")
    return {"text": missing_stock_text_fa(new, file_label=file_label, actor=actor, has_excel=bool(xlsx)),
            "xlsx": str(xlsx) if xlsx else None, "items": new}
